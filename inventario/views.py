# inventario/views.py
from django.views.generic import ListView, DetailView, CreateView, UpdateView, DeleteView, TemplateView
from django.urls import reverse_lazy
from .models import Producto, Movimiento, GuiaSalida, Almacen, Proveedor, UnidadMedida, NotaDespacho, DetalleNotaDespacho, OrdenSalida, DetalleOrdenSalida, ActaRecepcion, DetalleActaRecepcion, Cliente, Cotizacion, DetalleCotizacion, PagoProveedor # Importa los nuevos modelos
from django.forms import ModelForm, inlineformset_factory # Import inlineformset_factory
from django import forms # Importar forms para widgets
from django.db.models import Sum, F, Q # Importa Sum, F y Q para consultas avanzadas
import django.db.models as models
from django.http import HttpResponse, Http404, JsonResponse, HttpResponseRedirect
from django.template.loader import get_template, render_to_string
from xhtml2pdf import pisa
import json
import openpyxl
from openpyxl.styles import Font, Border, Side, Alignment
from datetime import datetime, date
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt
from django.db import transaction
from django.shortcuts import render, get_object_or_404, redirect
from io import BytesIO
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, A4
from .models import NotaDespacho # Or from your_app_name.models import NotaDespacho
from weasyprint import HTML
#from .forms import MovimientoForm

# Para códigos de barras
import barcode
from barcode.writer import ImageWriter
import base64

# Importar para manejar rutas de archivos estáticos
#from django.contrib.staticfiles.storage import staticfiles_storage
from django.conf import settings
import os
from django.forms.models import model_to_dict # Added this import for model_to_dict

# Importaciones para autenticación
from django.contrib.auth.views import LoginView, LogoutView
from django.contrib.auth.mixins import LoginRequiredMixin # Para proteger vistas


# --- Formularios ---
class ProductoForm(ModelForm):
    class Meta:
        model = Producto
        fields = ['nombre', 'descripcion', 'sku', 'cantidad', 'precio', 'unidad_medida', 'proveedor', 'almacen', 'stock_minimo', 'marca', 'modelo', 'foto']

class MovimientoForm(ModelForm):
    class Meta:
        model = Movimiento
        fields = ['producto', 'tipo', 'cantidad', 'descripcion', 'guia_salida']

class GuiaSalidaForm(ModelForm):
    class Meta:
        model = GuiaSalida
        fields = ['destinatario', 'nota_despacho', 'acta_recepcion']

class NotaDespachoForm(ModelForm):
    class Meta:
        model = NotaDespacho
        # Corrected to use 'orden_salida_referencia'
        fields = ['cliente', 'orden_salida_referencia', 'proveedor','orden_asociada', 'observaciones', 'nombre_conductor', 'ci_conductor', 'tipo_vehiculo', 'color_vehiculo', 'placa_vehiculo']

class DetalleNotaDespachoForm(ModelForm):
    class Meta:
        model = DetalleNotaDespacho
        # Added detalle_cotizacion_origen
        fields = ['producto', 'detalle_cotizacion_origen', 'nombre_producto_despacho', 'marca', 'modelo', 'cantidad']

class OrdenSalidaForm(ModelForm):
    class Meta:
        model = OrdenSalida
        fields = ['cliente', 'cotizacion_origen']

class DetalleOrdenSalidaForm(ModelForm):
    class Meta:
        model = DetalleOrdenSalida
        fields = ['producto', 'cantidad', 'precio_unitario', 'subtotal']

class ActaRecepcionForm(ModelForm):
    class Meta:
        model = ActaRecepcion
        fields = ['proveedor', 'orden_asociada', 'ubicacion_almacen', 'estado_mercancia', 'observaciones', 'nombre_conductor', 'ci_conductor', 'tipo_vehiculo', 'color_vehiculo', 'placa_vehiculo', 'estado'] # Added 'estado'

class DetalleActaRecepcionForm(ModelForm):
    class Meta:
        model = DetalleActaRecepcion
        fields = ['producto', 'nombre_producto_recepcion', 'marca', 'modelo', 'cantidad']

class ClienteForm(ModelForm):
    class Meta:
        model = Cliente
        fields = ['nombre', 'cedula_rif', 'direccion', 'telefono', 'email']

class CotizacionForm(ModelForm):
    class Meta:
        model = Cotizacion
        fields = ['cliente', 'fecha_validez', 'estado', 'observaciones']
        widgets = {
            'fecha_validez': forms.DateInput(attrs={'type': 'date'}, format='%Y-%m-%d'),
        }

class DetalleCotizacionForm(ModelForm):
    class Meta:
        model = DetalleCotizacion
        fields = ['producto', 'nombre_producto_cotizado', 'marca', 'modelo', 'cantidad', 'precio_unitario']

# Nuevo Formulario: PagoProveedorForm
class PagoProveedorForm(ModelForm):
    class Meta:
        model = PagoProveedor
        fields = ['proveedor', 'monto', 'fecha_pago', 'descripcion']
        widgets = {
            'fecha_pago': forms.DateInput(attrs={'type': 'date'}, format='%Y-%m-%d'),
        }


# --- Vistas de Autenticación ---
class CustomLoginView(LoginView):
    template_name = 'inventario/login.html'
    redirect_authenticated_user = True # Redirige a los usuarios ya autenticados

    def get_success_url(self):
        return reverse_lazy('inventario:dashboard') # Redirige al dashboard después del login

class CustomLogoutView(LogoutView):
    next_page = reverse_lazy('inventario:login') # Redirige a la página de login después del logout


# --- Vistas para Dashboard ---
class DashboardView(LoginRequiredMixin, TemplateView): # Protegida con LoginRequiredMixin
    template_name = 'inventario/dashboard.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['total_productos'] = Producto.objects.count()
        total_valor = 0
        for p in Producto.objects.all():
            total_valor += p.cantidad * p.precio
        context['stock_total_valor'] = total_valor
        
        context['ultimos_movimientos'] = Movimiento.objects.order_by('-fecha')[:5]
        context['productos_bajo_stock'] = Producto.objects.filter(cantidad__lt=models.F('stock_minimo')).order_by('cantidad')[:5]
        
        context['total_cotizaciones_abiertas'] = Cotizacion.objects.filter(estado='abierta').count()
        context['total_ordenes_salida'] = OrdenSalida.objects.count()
        context['ultimas_cotizaciones'] = Cotizacion.objects.order_by('-fecha_creacion')[:5]
        context['ultimas_actas_recepcion'] = ActaRecepcion.objects.order_by('-fecha_recepcion')[:5]

        return context

# --- Vistas para Producto ---
class ProductoListView(LoginRequiredMixin, ListView):
    model = Producto
    template_name = 'inventario/product_list.html'
    context_object_name = 'products'

class ProductoDetailView(LoginRequiredMixin, DetailView):
    model = Producto
    template_name = 'inventario/product_detail.html'
    context_object_name = 'products'

class ProductoCreateView(LoginRequiredMixin, CreateView):
    model = Producto
    form_class = ProductoForm
    template_name = 'inventario/product_form.html'
    success_url = reverse_lazy('inventario:product_list')

class ProductoUpdateView(LoginRequiredMixin, UpdateView):
    model = Producto
    form_class = ProductoForm
    template_name = 'inventario/product_form.html'
    success_url = reverse_lazy('inventario:product_list')

class ProductoDeleteView(LoginRequiredMixin, DeleteView):
    model = Producto
    template_name = 'inventario/product_confirm_delete.html'
    success_url = reverse_lazy('inventario:product_list')

# --- Funciones de Exportación/Importación para Productos ---

def export_productos_excel(request):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="productos_inventario.xlsx"'

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Productos de Inventario"

    image_path = os.path.join(settings.BASE_DIR, 'inventario', 'static', 'img', 'encabezado_oficial.jpg')
    
    if os.path.exists(image_path):
        print(f"Advertencia: La imagen del membrete se encontró en {image_path}, pero la inserción en Excel requiere 'openpyxl.drawing.image.Image'.")
    else:
        print(f"Advertencia: La imagen del membrete no se encontró en {image_path}")

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = openpyxl.styles.PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(left=Side(style='thin'), 
                         right=Side(style='thin'), 
                         top=Side(style='thin'), 
                         bottom=Side(style='thin'))

    headers = ["ID", "Nombre", "Descripción", "SKU", "Cantidad", "Precio", "Unidad de Medida", "Proveedor", "Almacén", "Stock Mínimo", "Marca", "Modelo", "Fecha Creación", "Fecha Actualización"]
    sheet.append(headers)

    for col_num, cell in enumerate(sheet[sheet.max_row], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        sheet.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 20

    productos = Producto.objects.all().order_by('nombre')
    for producto in productos:
        row_data = [
            producto.id,
            producto.nombre,
            producto.descripcion,
            producto.sku,
            producto.cantidad,
            float(producto.precio),
            producto.unidad_medida.abreviatura if producto.unidad_medida else '',
            producto.proveedor.nombre if producto.proveedor else '',
            producto.almacen.nombre if producto.almacen else '',
            producto.stock_minimo,
            producto.marca,
            producto.modelo,
            producto.fecha_creacion.strftime("%Y-%m-%d %H:%M:%S"),
            producto.fecha_actualizacion.strftime("%Y-%m-%d %H:%M:%S"),
        ]
        sheet.append(row_data)
        for cell in sheet[sheet.max_row]:
            cell.border = thin_border

    workbook.save(response)
    return response

def export_productos_pdf(request):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="productos_inventario.pdf"'

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()

    elements = []

    elements.append(Paragraph("Reporte de Productos de Inventario - ENVEPRO", styles['h1']))
    elements.append(Paragraph("Fecha del Reporte: " + timezone.now().strftime("%d/%m/%Y %H:%M"), styles['h3']))
    elements.append(Paragraph("<br/>", styles['Normal']))

    data = []
    data.append(["ID", "Nombre", "SKU", "Cantidad", "Precio", "Marca", "Modelo", "U. Medida", "Proveedor"])

    productos = Producto.objects.all().order_by('nombre')
    for producto in productos:
        row_data = [
            str(producto.id),
            producto.nombre,
            producto.sku,
            str(producto.cantidad),
            f"${producto.precio}",
            producto.marca,
            producto.modelo,
            producto.unidad_medida.abreviatura if producto.unidad_medida else 'N/A',
            producto.proveedor.nombre if producto.proveedor else 'N/A',
        ]
        data.append(row_data)

    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#4F81BD")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('BOX', (0, 0), (-1, -1), 1, colors.black),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('LEFTPADDING', (0,0), (-1,-1), 3),
        ('RIGHTPADDING', (0,0), (-1,-1), 3),
        ('WORDWRAP', (0,1), (-1,-1), True),
    ]))

    col_widths = [0.5*inch, 1.5*inch, 1*inch, 0.7*inch, 0.7*inch, 1*inch, 1*inch, 0.7*inch, 1*inch]
    table._argW = col_widths

    elements.append(table)
    doc.build(elements)

    response.write(buffer.getvalue())
    buffer.close()
    return response

def import_productos_excel(request):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    if request.method == 'POST':
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            return render(request, 'inventario/excel_upload.html', {'error_message': 'No se seleccionó ningún archivo.'})

        if not excel_file.name.endswith('.xlsx'):
            return render(request, 'inventario/excel_upload.html', {'error_message': 'Por favor, sube un archivo .xlsx válido.'})

        try:
            workbook = openpyxl.load_workbook(excel_file)
            sheet = workbook.active
            imported_count = 0
            updated_count = 0
            errors = []

            header = [cell.value for cell in sheet[1]]
            
            required_cols = ['Nombre', 'SKU', 'Cantidad', 'Precio']
            col_indices = {}
            for col_name in required_cols:
                try:
                    col_indices[col_name] = header.index(col_name)
                except ValueError:
                    return render(request, 'inventario/excel_upload.html', {'error_message': f'Falta la columna obligatoria "{col_name}" en el encabezado del archivo Excel.'})
            
            optional_cols = {
                'Descripción': 'descripcion',
                'Unidad de Medida': 'unidad_medida',
                'Proveedor': 'proveedor',
                'Almacén': 'almacen',
                'Stock Mínimo': 'stock_minimo',
                'Marca': 'marca',
                'Modelo': 'modelo',
            }
            for col_name, field_name in optional_cols.items():
                if col_name in header:
                    col_indices[col_name] = header.index(col_name)
                else:
                    col_indices[col_name] = -1

            for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                try:
                    nombre = row[col_indices['Nombre']].value
                    sku = row[col_indices['SKU']].value
                    cantidad = row[col_indices['Cantidad']].value
                    precio = row[col_indices['Precio']].value

                    if not nombre or not sku or cantidad is None or precio is None:
                        errors.append(f"Fila {row_idx}: Datos incompletos (Nombre, SKU, Cantidad o Precio son obligatorios).")
                        continue
                    
                    try:
                        cantidad = int(cantidad)
                        if cantidad < 0:
                            errors.append(f"Fila {row_idx}: Cantidad debe ser un número no negativo.")
                            continue
                    except (ValueError, TypeError):
                        errors.append(f"Fila {row_idx}: Cantidad inválida '{cantidad}'. Debe ser un número entero.")
                        continue
                    
                    try:
                        precio = float(precio)
                        if precio < 0:
                            errors.append(f"Fila {row_idx}: Precio debe ser un número no negativo.")
                            continue
                    except (ValueError, TypeError):
                        errors.append(f"Fila {row_idx}: Precio inválido '{precio}'. Debe ser un número.")
                        continue

                    defaults = {
                        'nombre': nombre,
                        'cantidad': cantidad,
                        'precio': precio,
                    }
                    
                    if col_indices['Descripción'] != -1:
                        defaults['descripcion'] = row[col_indices['Descripción']].value or ''
                    if col_indices['Stock Mínimo'] != -1:
                        try:
                            defaults['stock_minimo'] = int(row[col_indices['Stock Mínimo']].value) if row[col_indices['Stock Mínimo']].value is not None else 0
                        except (ValueError, TypeError):
                            errors.append(f"Fila {row_idx}: Stock Mínimo inválido. Se usará 0.")
                            defaults['stock_minimo'] = 0
                    if col_indices['Marca'] != -1:
                        defaults['marca'] = row[col_indices['Marca']].value or ''
                    if col_indices['Modelo'] != -1:
                        defaults['modelo'] = row[col_indices['Modelo']].value or ''

                    # Manejo de FKs
                    unidad_medida_nombre = row[col_indices['Unidad de Medida']].value if col_indices['Unidad de Medida'] != -1 else ''
                    if unidad_medida_nombre:
                        unidad_medida, _ = UnidadMedida.objects.get_or_create(nombre=unidad_medida_nombre, defaults={'abreviatura': unidad_medida_nombre[:10]})
                        defaults['unidad_medida'] = unidad_medida
                    else:
                        defaults['unidad_medida'] = None

                    proveedor_nombre = row[col_indices['Proveedor']].value if col_indices['Proveedor'] != -1 else ''
                    if proveedor_nombre:
                        proveedor, _ = Proveedor.objects.get_or_create(nombre=proveedor_nombre)
                        defaults['proveedor'] = proveedor
                    else:
                        defaults['proveedor'] = None

                    almacen_nombre = row[col_indices['Almacén']].value if col_indices['Almacén'] != -1 else ''
                    if almacen_nombre:
                        almacen, _ = Almacen.objects.get_or_create(nombre=almacen_nombre)
                        defaults['almacen'] = almacen
                    else:
                        defaults['almacen'] = None

                    producto, created = Producto.objects.update_or_create(
                        sku=sku,
                        defaults=defaults
                    )
                    if created:
                        imported_count += 1
                    else:
                        updated_count += 1

                except Exception as e:
                    errors.append(f"Fila {row_idx}: Error inesperado - {e}")
            
            if errors:
                return render(request, 'inventario/excel_upload.html', {
                    'success_message': f"Se importaron {imported_count} productos nuevos y se actualizaron {updated_count}. Sin embargo, se encontraron los siguientes errores:",
                    'errors': errors,
                    'is_product_import': True
                })
            else:
                return render(request, 'inventario/excel_upload.html', {'success_message': f"Se importaron {imported_count} productos nuevos y se actualizaron {updated_count} con éxito.", 'is_product_import': True})

        except Exception as e:
            return render(request, 'inventario/excel_upload.html', {'error_message': f"Error al procesar el archivo Excel: {e}", 'is_product_import': True})
    
    return render(request, 'inventario/excel_upload.html', {'is_product_import': True})


def generate_barcode_image(sku):
    """Genera una imagen de código de barras (EAN13 o Code128) para un SKU dado."""
    try:
        # Intentar generar EAN13 si el SKU es numérico y tiene longitud adecuada (12 o 13 dígitos)
        if sku and isinstance(sku, str) and sku.isdigit() and (len(sku) == 12 or len(sku) == 13):
            ean = barcode.EAN13(sku[:12], writer=ImageWriter()) # EAN13 usa 12 dígitos + 1 checksum
            fp = BytesIO()
            ean.write(fp)
            return base64.b64encode(fp.getvalue()).decode('utf-8')
        else:
            # Fallback a Code128 para SKUs no numéricos o de longitud arbitraria
            code128 = barcode.Code128(str(sku), writer=ImageWriter()) # Asegurar que SKU sea string
            fp = BytesIO()
            code128.write(fp)
            return base64.b64encode(fp.getvalue()).decode('utf-8')
    except Exception as e:
        print(f"Error generating barcode for SKU {sku}: {e}")
        return None

def export_product_labels_pdf(request):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    """
    Genera un PDF con etiquetas de productos, incluyendo nombre, SKU y código de barras,
    utilizando xhtml2pdf para renderizar la plantilla HTML.
    """
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="etiquetas_productos.pdf"'

    product_id = request.GET.get('product_id')
    if product_id:
        productos_a_etiquetar = Producto.objects.filter(pk=product_id)
    else:
        productos_a_etiquetar = Producto.objects.all().order_by('nombre')
    
    labels_data = []
    for producto in productos_a_etiquetar:
        barcode_base64 = generate_barcode_image(producto.sku)
        labels_data.append({
            'nombre': producto.nombre,
            'sku': producto.sku,
            'marca': producto.marca,
            'modelo': producto.modelo,
            'barcode_image': f"data:image/png;base64,{barcode_base64}" if barcode_base64 else None
        })
    
    products_per_row = 3
    labels_rows = []
    for i in range(0, len(labels_data), products_per_row):
        row = labels_data[i:i + products_per_row]
        while len(row) < products_per_row:
            row.append(None)
        labels_rows.append(row)
    
    context = {
        'labels_rows': labels_rows,
        'current_date': timezone.now().strftime("%d/%m/%Y"),
    }

    template_path = 'inventario/pdf/product_labels_pdf.html'
    html = render_to_string(template_path, context)

    pisa_status = pisa.CreatePDF(
        html,
        dest=response)

    if pisa_status.err:
        return HttpResponse('Error al generar el PDF de etiquetas: %s' % pisa_status.err, status=500)
    return response


# --- Vistas para Movimiento ---
class MovimientoListView(LoginRequiredMixin, ListView):
    model = Movimiento
    template_name = 'inventario/movimiento_list.html'
    context_object_name = 'movimientos'
    ordering = ['-fecha']

def export_movimientos_excel(request):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="movimientos_inventario.xlsx"'

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Movimientos de Inventario"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = openpyxl.styles.PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(left=Side(style='thin'), 
                         right=Side(style='thin'), 
                         top=Side(style='thin'), 
                         bottom=Side(style='thin'))

    headers = ["ID", "Producto", "SKU Producto", "Tipo", "Cantidad", "Fecha", "Descripción", "Guía de Salida"]
    sheet.append(headers)

    for col_num, cell in enumerate(sheet[1], 1):
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
        sheet.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 20

    movimientos = Movimiento.objects.all().order_by('fecha')
    for movimiento in movimientos:
        row_data = [
            movimiento.id,
            movimiento.producto.nombre if movimiento.producto else 'N/A',
            movimiento.producto.sku if movimiento.producto else 'N/A',
            movimiento.get_tipo_display(),
            movimiento.cantidad,
            movimiento.fecha.strftime("%Y-%m-%d %H:%M:%S"),
            movimiento.descripcion,
            movimiento.guia_salida.destinatario if movimiento.guia_salida else 'N/A',
        ]
        sheet.append(row_data)
        for cell in sheet[sheet.max_row]:
            cell.border = thin_border

    workbook.save(response)
    return response

def export_movimientos_pdf(request):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="movimientos_inventario.pdf"'

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    styles = getSampleStyleSheet()

    elements = []

    elements.append(Paragraph("Reporte de Movimientos de Inventario - ENVEPRO", styles['h1']))
    elements.append(Paragraph("Fecha del Reporte: " + timezone.now().strftime("%d/%m/%Y %H:%M"), styles['h3']))
    elements.append(Paragraph("<br/>", styles['Normal']))

    data = []
    data.append(["ID", "Producto", "SKU", "Tipo", "Cantidad", "Fecha", "Descripción"])

    movimientos = Movimiento.objects.all().order_by('fecha')
    for movimiento in movimientos:
        data.append([
            str(movimiento.id),
            movimiento.producto.nombre if movimiento.producto else 'N/A',
            movimiento.producto.sku if movimiento.producto else 'N/A',
            movimiento.get_tipo_display(),
            str(movimiento.cantidad),
            movimiento.fecha.strftime("%d/%m/%Y %H:%M"),
            movimiento.descripcion,
        ])

    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#4F81BD")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('BOX', (0, 0), (-1, -1), 1, colors.black),
        ('FONTSIZE', (0,0), (-1,-1), 8),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('LEFTPADDING', (0,0), (-1,-1), 3),
        ('RIGHTPADDING', (0,0), (-1,-1), 3),
        ('WORDWRAP', (0,1), (-1,-1), True),
    ]))

    col_widths = [0.5*inch, 1.5*inch, 1*inch, 0.7*inch, 0.7*inch, 1.5*inch, 2*inch]
    table._argW = col_widths

    elements.append(table)
    doc.build(elements)

    response.write(buffer.getvalue())
    buffer.close()
    return response

def import_movimientos_excel(request):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    if request.method == 'POST':
        excel_file = request.FILES.get('excel_file')
        if not excel_file:
            return render(request, 'inventario/excel_upload.html', {'error_message': 'No se seleccionó ningún archivo.'})

        if not excel_file.name.endswith('.xlsx'):
            return render(request, 'inventario/excel_upload.html', {'error_message': 'Por favor, sube un archivo .xlsx válido.'})

        try:
            workbook = openpyxl.load_workbook(excel_file)
            sheet = workbook.active
            imported_count = 0
            errors = []

            header = [cell.value for cell in sheet[1]]
            
            try:
                col_sku = header.index('SKU del Producto')
                col_tipo = header.index('Tipo de Movimiento')
                col_cantidad = header.index('Cantidad')
                col_descripcion = header.index('Descripción') if 'Descripción' in header else -1
                col_fecha = header.index('Fecha') if 'Fecha' in header else -1
            except ValueError as e:
                return render(request, 'inventario/excel_upload.html', {'error_message': f'Faltan columnas obligatorias en el encabezado del archivo Excel. Asegúrate de incluir "SKU del Producto", "Tipo de Movimiento" y "Cantidad". Detalle: {e}'})


            for row_idx, row in enumerate(sheet.iter_rows(min_row=2), start=2):
                try:
                    sku = row[col_sku].value
                    tipo = str(row[col_tipo].value).lower()
                    cantidad = row[col_cantidad].value
                    descripcion = row[col_descripcion].value if col_descripcion != -1 else ''
                    fecha_val = row[col_fecha].value if col_fecha != -1 else None

                    if not sku or not tipo or not cantidad:
                        errors.append(f"Fila {row_idx}: Datos incompletos (SKU, Tipo o Cantidad son obligatorios).")
                        continue
                    
                    if tipo not in ['entrada', 'salida']:
                        errors.append(f"Fila {row_idx}: Tipo de movimiento inválido '{tipo}'. Debe ser 'entrada' o 'salida'.")
                        continue

                    try:
                        cantidad = int(cantidad)
                        if cantidad <= 0:
                            errors.append(f"Fila {row_idx}: Cantidad debe ser un número positivo.")
                            continue
                    except (ValueError, TypeError):
                        errors.append(f"Fila {row_idx}: Cantidad inválida '{cantidad}'. Debe ser un número entero.")
                        continue

                    try:
                        producto = Producto.objects.get(sku=sku)
                    except Producto.DoesNotExist:
                        errors.append(f"Fila {row_idx}: Producto con SKU '{sku}' no encontrado.")
                        continue

                    movimiento_fecha = timezone.now()
                    if fecha_val:
                        try:
                            if isinstance(fecha_val, datetime):
                                movimiento_fecha = timezone.make_aware(fecha_val)
                            elif isinstance(fecha_val, date):
                                movimiento_fecha = timezone.make_aware(datetime.combine(fecha_val, datetime.min.time()))
                            else:
                                parsed_date = None
                                formats_to_try = [
                                    '%Y-%m-%d %H:%M:%S',
                                    '%Y-%m-%d %H:%M',
                                    '%Y-%m-%d',
                                    '%d/%m/%Y %H:%M:%S',
                                    '%d/%m/%Y %H:%M',
                                    '%d/%m/%Y',
                                ]
                                for fmt in formats_to_try:
                                    try:
                                        parsed_date = datetime.strptime(str(fecha_val), fmt)
                                        break
                                    except ValueError:
                                        pass
                                
                                if parsed_date:
                                    movimiento_fecha = timezone.make_aware(parsed_date)
                                else:
                                    errors.append(f"Fila {row_idx}: Formato de fecha inválido '{fecha_val}'. Use IHDA-MM-DD HH:MM:SS u otros formatos comunes.")
                        except Exception as e:
                            errors.append(f"Fila {row_idx}: Error al procesar fecha '{fecha_val}': {e}.")
                    
                    Movimiento.objects.create(
                        producto=producto,
                        tipo=tipo,
                        cantidad=cantidad,
                        descripcion=descripcion,
                        fecha=movimiento_fecha
                    )
                    imported_count += 1

                except Exception as e:
                    errors.append(f"Fila {row_idx}: Error inesperado - {e}")
            
            if errors:
                return render(request, 'inventario/excel_upload.html', {
                    'success_message': f"Se importaron {imported_count} movimientos con éxito. Sin embargo, se encontraron los siguientes errores:",
                    'errors': errors,
                    'is_product_import': False
                })
            else:
                return render(request, 'inventario/excel_upload.html', {'success_message': f"Se importaron {imported_count} movimientos con éxito.", 'is_product_import': False})

        except Exception as e:
            return render(request, 'inventario/excel_upload.html', {'error_message': f"Error al procesar el archivo Excel: {e}", 'is_product_import': False})
    
    return render(request, 'inventario/excel_upload.html', {'is_product_import': False})


class MovimientoCreateView(LoginRequiredMixin, CreateView):
    model = Movimiento
    form_class = MovimientoForm
    template_name = 'inventario/movimiento_form.html'
    success_url = reverse_lazy('inventario:movimiento_list')

class MovimientoDetailView(LoginRequiredMixin, DetailView):
    model = Movimiento
    template_name = 'inventario/movimiento_detail.html'
    context_object_name = 'movimiento'

class MovimientoUpdateView(UpdateView):
    model = Movimiento
    form_class = MovimientoForm
    template_name = 'inventario/movimiento_form.html'
    success_url = reverse_lazy('inventario:movimiento_list')

    def form_valid(self, form):
        # Obtenemos el objeto Movimiento original antes de que el formulario lo modifique.
        # Esto es crucial para calcular el impacto del cambio en el stock.
        original_movimiento = self.get_object()
        
        # Obtenemos los datos del formulario (los nuevos valores)
        new_cantidad = form.cleaned_data['cantidad']
        new_tipo = form.cleaned_data['tipo']
        new_producto = form.cleaned_data['producto'] # El producto seleccionado en el formulario

        # Obtenemos el producto asociado al movimiento original
        original_product = original_movimiento.producto
        
        # --- Paso 1: Calcular el cambio de stock al "revertir" el movimiento original ---
        # Si el movimiento original era una 'entrada', al editarlo, esa cantidad "sale" del stock.
        # Si el movimiento original era una 'salida', al editarlo, esa cantidad "vuelve" al stock.
        stock_change_from_original = 0
        if original_movimiento.tipo == 'entrada':
            stock_change_from_original = -original_movimiento.cantidad # Revertir la entrada
        elif original_movimiento.tipo == 'salida':
            stock_change_from_original = original_movimiento.cantidad # Revertir la salida

        # --- Paso 2: Calcular el cambio de stock al "aplicar" el nuevo movimiento ---
        # Si el nuevo movimiento es una 'entrada', la cantidad "entra" al stock.
        # Si el nuevo movimiento es una 'salida', la cantidad "sale" del stock.
        stock_change_for_new = 0
        if new_tipo == 'entrada':
            stock_change_for_new = new_cantidad
        elif new_tipo == 'salida':
            stock_change_for_new = -new_cantidad

        # --- Paso 3: Realizar la verificación del stock hipotético ---
        # Consideramos dos escenarios: si el producto asociado al movimiento cambia o no.
        if original_product.pk != new_producto.pk:
            # Escenario 1: El producto asociado al movimiento está cambiando.
            # Necesitamos verificar el stock de AMBOS productos.

            # a) Verificar el stock del producto ORIGINAL después de revertir el movimiento.
            hypothetical_old_product_stock = original_product.cantidad + stock_change_from_original
            if hypothetical_old_product_stock < 0:
                form.add_error(None, f"No se puede actualizar el movimiento. Revertir el movimiento original de '{original_product.nombre}' resultaría en stock negativo ({hypothetical_old_product_stock}).")
                return self.form_invalid(form)

            # b) Verificar el stock del producto NUEVO después de aplicar el nuevo movimiento.
            hypothetical_new_product_stock = new_producto.cantidad + stock_change_for_new
            if hypothetical_new_product_stock < 0:
                form.add_error(None, f"No se puede actualizar el movimiento. La cantidad de '{new_producto.nombre}' sería negativa con el nuevo movimiento ({hypothetical_new_product_stock}).")
                return self.form_invalid(form)
        else:
            # Escenario 2: El producto asociado al movimiento NO está cambiando.
            # Calculamos el cambio neto en el stock del mismo producto.
            # Stock actual del producto + (cambio por revertir viejo) + (cambio por aplicar nuevo)
            net_stock_after_update = original_product.cantidad + stock_change_from_original + stock_change_for_new
            if net_stock_after_update < 0:
                form.add_error(None, f"No se puede actualizar el movimiento. La cantidad de '{original_product.nombre}' sería negativa ({net_stock_after_update}).")
                return self.form_invalid(form)

        # Si todas las verificaciones pasan, procedemos a guardar el formulario.
        # Asumimos que las señales (pre_save/post_save) del modelo Movimiento
        # se encargarán de la actualización real del stock del producto.
        return super().form_valid(form)

class MovimientoDeleteView(LoginRequiredMixin, DeleteView):
    model = Movimiento
    template_name = 'inventario/movimiento_confirm_delete.html'
    success_url = reverse_lazy('inventario:movimiento_list')

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        movimiento = self.object
        product = movimiento.producto

        if product:
            if movimiento.tipo == 'entrada':
                # Si se elimina un movimiento de 'entrada', la cantidad del producto disminuye.
                # Verificamos si esto resultaría en una cantidad negativa.
                new_quantity = product.cantidad - movimiento.cantidad
                if new_quantity < 0:
                    # Si la cantidad sería negativa, impedimos la eliminación
                    # y mostramos un mensaje de error.
                    return self.render_to_response(self.get_context_data(
                        object=self.object,
                        error_message=f"No se puede eliminar este movimiento de entrada. La cantidad de '{product.nombre}' sería negativa ({new_quantity})."
                    ))
            elif movimiento.tipo == 'salida':
                # Si se elimina un movimiento de 'salida', la cantidad del producto aumenta.
                # No necesitamos una verificación de negativo aquí ya que la cantidad aumenta.
                new_quantity = product.cantidad + movimiento.cantidad
            
            # Si la verificación pasa (o no es aplicable para 'salida'),
            # procedemos con la eliminación dentro de una transacción atómica.
            # Esto asegura que si algo falla durante la eliminación o las actualizaciones
            # de señales (si las tienes), la base de datos se revierta a su estado anterior.
            with transaction.atomic():
                # La señal post_delete del modelo Movimiento (si la tienes configurada)
                # debería manejar la actualización real del stock del producto.
                # Si no tienes una señal, tendrías que actualizar product.cantidad aquí:
                # product.cantidad = new_quantity
                # product.save()
                return super().post(request, *args, **kwargs)
        
        # Si no hay producto asociado, o si no entra en las condiciones anteriores,
        # simplemente se procede con la eliminación (esto podría indicar un caso de borde
        # o un movimiento sin producto, lo cual debería ser manejado por tu lógica de negocio).
        return super().post(request, *args, **kwargs)
# --- Vistas para GuiaSalida ---
class GuiaSalidaListView(LoginRequiredMixin, ListView):
    model = GuiaSalida
    template_name = 'inventario/guia_salida_list.html'
    context_object_name = 'guias_salida'
    ordering = ['-fecha_creacion']

class GuiaSalidaCreateView(LoginRequiredMixin, CreateView):
    model = GuiaSalida
    form_class = GuiaSalidaForm
    template_name = 'inventario/guia_salida_form.html'
    success_url = reverse_lazy('inventario:guia_salida_list')

class GuiaSalidaDetailView(LoginRequiredMixin, DetailView):
    model = GuiaSalida
    template_name = 'inventario/guia_salida_detail.html'
    context_object_name = 'guia_salida'

class GuiaSalidaUpdateView(LoginRequiredMixin, UpdateView):
    model = GuiaSalida
    form_class = GuiaSalidaForm
    template_name = 'inventario/guia_salida_form.html'
    success_url = reverse_lazy('inventario:guia_salida_list')

class GuiaSalidaDeleteView(LoginRequiredMixin, DeleteView):
    model = GuiaSalida
    template_name = 'inventario/guia_salida_confirm_delete.html'
    success_url = reverse_lazy('inventario:guia_salida_list')

# Define the formset for DetalleNotaDespacho
DetalleNotaDespachoFormSet = inlineformset_factory(
    NotaDespacho,
    DetalleNotaDespacho,
    fields=('producto', 'cantidad', 'detalle_cotizacion_origen', 'nombre_producto_despacho', 'marca', 'modelo'),
    form=DetalleNotaDespachoForm,
    extra=1,
    can_delete=True
)


class NotaDespachoListView(LoginRequiredMixin, ListView):
    model = NotaDespacho
    template_name = 'inventario/nota_despacho_list.html'
    context_object_name = 'notas_despacho'
    ordering = ['-fecha_despacho']


class NotaDespachoDetailView(LoginRequiredMixin, DetailView):
    model = NotaDespacho
    template_name = 'inventario/nota_despacho_detail.html'
    context_object_name = 'nota_despacho'


def crear_o_editar_nota_despacho(request, pk=None):
    nota_despacho = None
    if pk:
        nota_despacho = get_object_or_404(NotaDespacho, pk=pk)

    if request.method == 'POST':
        form = NotaDespachoForm(request.POST, instance=nota_despacho)
        formset = DetalleNotaDespachoFormSet(request.POST, instance=nota_despacho, prefix='detalles')
        
        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                nota_guardada = form.save(commit=False)
                if not pk:
                    nota_guardada.creado_por = request.user
                nota_guardada.save()
                
                old_details = {str(d.pk): d for d in DetalleNotaDespacho.objects.filter(nota_despacho=nota_guardada)} if nota_guardada.pk else {}
                formset.instance = nota_guardada
                saved_details = formset.save(commit=False)
                new_details_map = {str(d.pk): d for d in saved_details if d.pk}

                for form_data in formset.forms:
                    producto_instance = form_data.cleaned_data.get('producto')
                    product_id = producto_instance.pk if producto_instance else None
                    cantidad_despachada = form_data.cleaned_data.get('cantidad', 0)
                    detalle_cotizacion_origen = form_data.cleaned_data.get('detalle_cotizacion_origen')
                    detail_pk = form_data.cleaned_data.get('id')
                    
                    if form_data.cleaned_data.get('DELETE'):
                        if detail_pk and str(detail_pk) in old_details:
                            old_detail = old_details[str(detail_pk)]
                            if old_detail.producto:
                                old_detail.producto.cantidad += old_detail.cantidad
                                old_detail.producto.save()
                            old_detail.delete()
                    elif product_id and cantidad_despachada > 0:
                        if detail_pk and str(detail_pk) in old_details:
                            old_detail = old_details[str(detail_pk)]
                            if old_detail.producto:
                                diff = cantidad_despachada - old_detail.cantidad
                                old_detail.producto.cantidad -= diff
                                old_detail.producto.save()
                            old_detail.nombre_producto_despacho = form_data.cleaned_data.get('nombre_producto_despacho')
                            old_detail.marca = form_data.cleaned_data.get('marca')
                            old_detail.modelo = form_data.cleaned_data.get('modelo')
                            old_detail.cantidad = cantidad_despachada
                            old_detail.detalle_cotizacion_origen = detalle_cotizacion_origen
                            old_detail.save()
                        else:
                            new_detail = form_data.save(commit=False)
                            new_detail.nota_despacho = nota_guardada
                            new_detail.save()
                            
                            product = new_detail.producto
                            if product:
                                if product.cantidad < new_detail.cantidad:
                                    form.add_error(None, f"Stock insuficiente para {product.nombre}. Disponible: {product.cantidad}")
                                    transaction.set_rollback(True)
                                    cotizaciones = Cotizacion.objects.filter(estado='aprobada').order_by('-fecha_creacion')
                                    all_products_data = list(Producto.objects.all().values('id', 'nombre', 'sku', 'marca', 'modelo', 'cantidad', 'unidad_medida__abreviatura'))
                                    all_products_json = json.dumps(all_products_data)
                                    context = {
                                        'form': form,
                                        'detalle_formset': formset,
                                        'all_products_json': all_products_json,
                                        'cotizaciones': cotizaciones,
                                        'nota_despacho': nota_despacho,
                                    }
                                    return render(request, 'inventario/nota_despacho_form.html', context)
                                product.cantidad -= new_detail.cantidad
                                product.save()
            
            return redirect('inventario:nota_despacho_list')
        else:
            pass

    else:
        form = NotaDespachoForm(instance=nota_despacho)
        formset = DetalleNotaDespachoFormSet(instance=nota_despacho, prefix='detalles')

    cotizaciones = Cotizacion.objects.filter(estado='aprobada').order_by('-fecha_creacion')

    all_products_data = list(Producto.objects.all().values(
        'id', 'nombre', 'sku', 'marca', 'modelo', 'cantidad', 'unidad_medida__abreviatura'
    ))
    all_products_json = json.dumps(all_products_data)

    initial_detalles_data = []
    if nota_despacho:
        for detalle in nota_despacho.detalles.all():
            initial_detalles_data.append({
                'id': detalle.pk,
                'producto': detalle.producto.pk if detalle.producto else None,
                'nombre_producto_despacho': detalle.nombre_producto_despacho,
                'marca': detalle.marca,
                'modelo': detalle.modelo,
                'cantidad': float(detalle.cantidad),
                'detalle_cotizacion_origen': detalle.detalle_cotizacion_origen.pk if detalle.detalle_cotizacion_origen else None,
            })
    initial_detalles_json = json.dumps(initial_detalles_data)

    context = {
        'form': form,
        'detalle_formset': formset,
        'all_products_json': all_products_json,
        'initial_dispatch_items_json': initial_detalles_json,
        'cotizaciones': cotizaciones,
        'nota_despacho': nota_despacho,
    }
    return render(request, 'inventario/nota_despacho_form.html', context)


def crear_nota_despacho(request):
    return crear_o_editar_nota_despacho(request)


def editar_nota_despacho(request, pk):
    return crear_o_editar_nota_despacho(request, pk)


def get_productos_from_cotizacion(request, pk):
    try:
        cotizacion = Cotizacion.objects.get(pk=pk)
        detalles = DetalleCotizacion.objects.filter(cotizacion=cotizacion).select_related('producto')
        
        productos_json = []
        for detalle in detalles:
            productos_json.append({
                'id': detalle.producto.pk,
                'name': detalle.nombre_producto_cotizacion or detalle.producto.nombre,
                'brand': detalle.marca or detalle.producto.marca,
                'model': detalle.modelo or detalle.producto.modelo,
                'quantity': float(detalle.cantidad),
                'detalleCotizacionOrigenPk': detalle.pk,
            })
        return JsonResponse({'success': True, 'productos': productos_json})
    except Cotizacion.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'Cotización no encontrada'}, status=404)
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)}, status=500)


class NotaDespachoDeleteView(LoginRequiredMixin, DeleteView):
    model = NotaDespacho
    template_name = 'inventario/nota_despacho_confirm_delete.html'
    success_url = reverse_lazy('inventario:nota_despacho_list')

    def form_valid(self, form):
        with transaction.atomic():
            nota = self.get_object()
            for detalle in nota.detalles.all():
                if detalle.producto:
                    detalle.producto.cantidad += detalle.cantidad
                    detalle.producto.save()
            return super().form_valid(form)


def export_nota_despacho_pdf(request, pk):
    if not request.user.is_authenticated:
        return redirect('inventario:login')
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="nota_despacho_{pk}.pdf"'
    return response


def generate_nota_despacho_pdf(request, pk):
    nota_despacho = get_object_or_404(NotaDespacho.objects.select_related('proveedor'), pk=pk)
    detalles = nota_despacho.detalles.all()
    
    total_cantidad = sum(d.cantidad for d in detalles)
    
    min_rows = 10
    empty_rows_count = max(0, min_rows - len(detalles))
    empty_rows_range = range(empty_rows_count)

    context = {
        'nota_despacho': nota_despacho,
        'detalles': detalles,
        'total_cantidad': total_cantidad,
        'empty_rows_range': empty_rows_range,
        'cliente_data': nota_despacho.cliente.nombre,
        'proveedor_nombre': nota_despacho.proveedor.nombre if nota_despacho.proveedor else 'N/A',
        'nombre_conductor': nota_despacho.nombre_conductor,
        'ci_conductor': nota_despacho.ci_conductor,
        'tipo_vehiculo': nota_despacho.tipo_vehiculo,
        'color_vehiculo': nota_despacho.color_vehiculo,
        'placa_vehiculo': nota_despacho.placa_vehiculo,
        'logo_path': request.build_absolute_uri('/static/img/image_3fa1f6.png'),
    }

    template = get_template('inventario/nota_despacho_pdf.html')
    html_string = template.render(context)
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="nota_despacho_{nota_despacho.numero_despacho or nota_despacho.pk}.pdf"'

    HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf(response)
    return response

    # Renderiza la plantilla HTML a una cadena
    template = get_template('inventario/nota_despacho_pdf.html')
    html_string = template.render(context)
    # 🔍 LÍNEAS DE DEPURACIÓN
    print("Nota de Despacho:", nota_despacho)
    print("Cantidad de Detalles:", detalles.count())
    print("Primeros 500 caracteres del HTML generado:\n", html_string[:500])

    # Crea la respuesta HTTP con tipo de contenido PDF
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="nota_despacho_{nota_despacho.numero_despacho or nota_despacho.pk}.pdf"'

    # Convierte el HTML a PDF usando WeasyPrint
    HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf(response)
    return response


# --- Vistas para la Interfaz de Cotizaciones (anteriormente Mini POS) ---

class CotizacionesInterfaceView(LoginRequiredMixin, TemplateView):
    template_name = 'inventario/cotizaciones_interface.html' # Nuevo nombre de plantilla

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        # Obtener todos los productos disponibles para la cotización
        context['productos'] = Producto.objects.filter(cantidad__gt=0).order_by('nombre')
        context['clientes'] = Cliente.objects.all().order_by('nombre') # Pasar clientes para el selector
        return context

@csrf_exempt # ¡Advertencia! Esto es solo para desarrollo. En producción, usa CSRF tokens.
def create_cotizacion_from_interface(request):
    if not request.user.is_authenticated:
        return JsonResponse({'status': 'error', 'message': 'No autenticado.'}, status=401)
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            cliente_id = data.get('cliente_id') # Ahora esperamos un ID de cliente
            items = data.get('items', [])

            if not cliente_id:
                return JsonResponse({'status': 'error', 'message': 'Debe seleccionar un cliente.'}, status=400)
            
            try:
                cliente = Cliente.objects.get(id=cliente_id)
            except Cliente.DoesNotExist:
                return JsonResponse({'status': 'error', 'message': 'Cliente no encontrado.'}, status=400)

            if not items:
                return JsonResponse({'status': 'error', 'message': 'No hay ítems en la cotización.'}, status=400)

            total_cotizacion = 0
            detalles_para_cotizacion = []

            with transaction.atomic():
                # Validar y preparar ítems para la cotización
                for item_data in items:
                    product_id = item_data.get('producto_id')
                    cantidad_solicitada = item_data.get('cantidad')
                    precio_unitario = item_data.get('precio_unitario') # Precio puede ser el del producto o uno manual

                    if not product_id or not cantidad_solicitada or precio_unitario is None:
                        raise ValueError("Datos de ítem incompletos para la cotización.")
                    
                    try:
                        producto = Producto.objects.get(id=product_id)
                    except Producto.DoesNotExist:
                        raise ValueError(f"Producto con ID {product_id} no encontrado.")

                    if cantidad_solicitada <= 0:
                        raise ValueError(f"Cantidad inválida para {producto.nombre}.")
                    
                    # Usar el nombre, marca y modelo del producto si no se proporcionan explícitamente en el JSON
                    nombre_producto_cotizado = item_data.get('nombre_producto_cotizado', producto.nombre)
                    marca = item_data.get('marca', producto.marca)
                    modelo = item_data.get('modelo', producto.modelo)

                    subtotal_item = float(precio_unitario) * cantidad_solicitada
                    total_cotizacion += subtotal_item
                    detalles_para_cotizacion.append({
                        'producto': producto,
                        'cantidad': cantidad_solicitada,
                        'precio_unitario': float(precio_unitario),
                        'subtotal': subtotal_item,
                        'nombre_producto_cotizado': nombre_producto_cotizado,
                        'marca': marca,
                        'modelo': modelo
                    })
                
                # Crear la Cotización
                cotizacion = Cotizacion.objects.create(
                    cliente=cliente,
                    total_cotizado=total_cotizacion,
                    estado='abierta' # Estado inicial de una cotización
                )

                for item_data in detalles_para_cotizacion:
                    DetalleCotizacion.objects.create(
                        cotizacion=cotizacion,
                        producto=item_data['producto'],
                        cantidad=item_data['cantidad'],
                        precio_unitario=item_data['precio_unitario'],
                        subtotal=item_data['subtotal'],
                        nombre_producto_cotizado=item_data['nombre_producto_cotizado'],
                        marca=item_data['marca'],
                        modelo=item_data['modelo']
                    )

            return JsonResponse({'status': 'success', 'message': f'Cotización {cotizacion.numero_cotizacion} creada con éxito.', 'cotizacion_id': cotizacion.id})

        except ValueError as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': f'Error interno del servidor: {e}'}, status=500)
    
    return JsonResponse({'status': 'error', 'message': 'Método no permitido.'}, status=405)

@csrf_exempt
def get_cotizacion_details_api(request, pk):
    if not request.user.is_authenticated:
        return JsonResponse({'status': 'error', 'message': 'No autenticado.'}, status=401)
    """
    API endpoint to get details of a Cotizacion, including its DetalleCotizacion items.
    """
    try:
        cotizacion = Cotizacion.objects.get(pk=pk)
        detalles = cotizacion.detalles.all().values(
            'id', 'producto_id', 'nombre_producto_cotizado', 'marca', 'modelo', 'cantidad', 'cantidad_entregada', 'precio_unitario'
        )
        return JsonResponse({'status': 'success', 'detalles': list(detalles)})
    except Cotizacion.DoesNotExist:
        return JsonResponse({'status': 'error', 'message': 'Cotización no encontrada.'}, status=404)
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)

# --- Vistas para ActaRecepcion ---
# Define the formset for DetalleActaRecepcion
DetalleActaRecepcionFormSet = inlineformset_factory(
    ActaRecepcion,
    DetalleActaRecepcion,
    form=DetalleActaRecepcionForm,
    extra=0, # Changed to 0 as JS will manage adding rows
    can_delete=True
)

class ActaRecepcionListView(LoginRequiredMixin, ListView):
    model = ActaRecepcion
    template_name = 'inventario/acta_recepcion_list.html' # Nueva plantilla
    context_object_name = 'actas_recepcion'
    ordering = ['-fecha_recepcion']

class ActaRecepcionDetailView(LoginRequiredMixin, DetailView):
    model = ActaRecepcion
    template_name = 'inventario/acta_recepcion_detail.html' # Nueva plantilla
    context_object_name = 'acta_recepcion'

class ActaRecepcionCreateView(LoginRequiredMixin, CreateView):
    model = ActaRecepcion
    form_class = ActaRecepcionForm
    template_name = 'inventario/acta_recepcion_form.html' # Nueva plantilla
    success_url = reverse_lazy('inventario:acta_recepcion_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.POST:
            context['detalle_formset'] = DetalleActaRecepcionFormSet(self.request.POST, instance=self.object)
        else:
            context['detalle_formset'] = DetalleActaRecepcionFormSet(instance=self.object)
        
        context['all_products'] = list(Producto.objects.all().values('id', 'nombre', 'sku', 'marca', 'modelo', 'cantidad', 'unidad_medida__abreviatura'))
        
        # Serializar los datos iniciales del formset a JSON para JavaScript
        initial_detalles_data = []
        if self.object and self.object.pk: # Si estamos editando una Acta existente
            for detalle in self.object.detalles.all():
                initial_detalles_data.append({
                    'id': detalle.id,
                    'producto': detalle.producto.id if detalle.producto else None,
                    'nombre_producto_recepcion': detalle.nombre_producto_recepcion,
                    'marca': detalle.marca,
                    'modelo': detalle.modelo,
                    'cantidad': detalle.cantidad,
                })
        context['detalle_formset'].initial_json = json.dumps(initial_detalles_data)

        return context

    def form_valid(self, form):
        context = self.get_context_data()
        detalle_formset = context['detalle_formset']
        with transaction.atomic():
            self.object = form.save()
            if detalle_formset.is_valid():
                detalle_formset.instance = self.object
                detalle_formset.save()
            else:
                return self.form_invalid(form)
        return super().form_valid(form)

class ActaRecepcionUpdateView(LoginRequiredMixin, UpdateView):
    model = ActaRecepcion
    form_class = ActaRecepcionForm
    template_name = 'inventario/acta_recepcion_form.html' # Nueva plantilla
    success_url = reverse_lazy('inventario:acta_recepcion_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.POST:
            context['detalle_formset'] = DetalleActaRecepcionFormSet(self.request.POST, instance=self.object)
        else:
            context['detalle_formset'] = DetalleActaRecepcionFormSet(instance=self.object)
        
        context['all_products'] = list(Producto.objects.all().values('id', 'nombre', 'sku', 'marca', 'modelo', 'cantidad', 'unidad_medida__abreviatura'))
        
        # Serializar los datos iniciales del formset a JSON para JavaScript
        initial_detalles_data = []
        if self.object and self.object.pk: # Si estamos editando una Acta existente
            for detalle in self.object.detalles.all():
                initial_detalles_data.append({
                    'id': detalle.id,
                    'producto': detalle.producto.id if detalle.producto else None,
                    'nombre_producto_recepcion': detalle.nombre_producto_recepcion,
                    'marca': detalle.marca,
                    'modelo': detalle.modelo,
                    'cantidad': detalle.cantidad,
                })
        context['detalle_formset'].initial_json = json.dumps(initial_detalles_data)

        return context

    def form_valid(self, form):
        context = self.get_context_data()
        detalle_formset = context['detalle_formset']
        with transaction.atomic():
            self.object = form.save()
            if detalle_formset.is_valid():
                detalle_formset.instance = self.object
                detalle_formset.save()
            else:
                return self.form_invalid(form)
        return super().form_valid(form)

class ActaRecepcionDeleteView(LoginRequiredMixin, DeleteView):
    model = ActaRecepcion
    template_name = 'inventario/acta_recepcion_confirm_delete.html' # Nueva plantilla
    success_url = reverse_lazy('inventario:acta_recepcion_list')

def export_acta_recepcion_pdf(request, pk):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    """
    Genera un Acta de Recepción en formato PDF utilizando xhtml2pdf.
    """
    try:
        acta_recepcion = get_object_or_404(ActaRecepcion, pk=pk)
    except Http404:
        return HttpResponse("Acta de Recepción no encontrada.", status=404)

    detalles = acta_recepcion.detalles.all()
    # Calcular cuántas filas vacías necesitamos para llegar a 5 ítems (como en la plantilla)
    num_empty_rows = max(0, 5 - len(detalles))
    empty_rows_range = range(num_empty_rows)

    # --- Obtener la ruta absoluta de la imagen ---
    image_static_path = 'img/encabezado_oficial.jpg'
    absolute_image_path = None
    try:
        absolute_image_path = staticfiles_storage.path(image_static_path)
        print(f"Ruta absoluta de la imagen para PDF (Acta): {absolute_image_path}") # Debugging
        if not os.path.exists(absolute_image_path):
            print(f"ADVERTENCIA (Acta): La imagen no existe en la ruta absoluta: {absolute_image_path}")
            absolute_image_path = None
    except NotImplementedError:
        print("ADVERTENCIA (Acta): staticfiles_storage.path() no soportado. Intentando con STATIC_ROOT.")
        if hasattr(settings, 'STATIC_ROOT') and settings.STATIC_ROOT:
            absolute_image_path = os.path.join(settings.STATIC_ROOT, image_static_path)
            if not os.path.exists(absolute_image_path):
                print(f"ADVERTENCIA (Acta): La imagen no existe en STATIC_ROOT: {absolute_image_path}")
                absolute_image_path = None
        else:
            print("ADVERTENCIA (Acta): STATIC_ROOT no está configurado, no se puede obtener la ruta absoluta de la imagen.")
            absolute_image_path = None
    except Exception as e:
        print(f"Error al obtener la ruta absoluta de la imagen (Acta): {e}")
        absolute_image_path = None

    context = {
        'acta_recepcion': acta_recepcion,
        'detalles': detalles,
        'total_cantidad': sum(detalle.cantidad for detalle in detalles),
        'current_date': timezone.now().strftime("%d/%m/%Y"),
        'empty_rows_range': empty_rows_range,
        'header_image_path': absolute_image_path,
    }

    template_path = 'inventario/pdf/acta_recepcion_pdf.html'
    html = render_to_string(template_path, context)

    pisa_status = pisa.CreatePDF(
        html,
        dest=response)

    if pisa_status.err:
        return HttpResponse('Error al generar el PDF: %s' % pisa_status.err, status=500)
    return response

class OrdenSalidaDetailView(LoginRequiredMixin, DetailView):
    model = OrdenSalida
    template_name = 'inventario/orden_salida_detail.html'
    context_object_name = 'orden_salida'

# --- Vistas para Clientes ---
class ClienteListView(LoginRequiredMixin, ListView):
    model = Cliente
    template_name = 'inventario/cliente_list.html'
    context_object_name = 'clientes'
    ordering = ['nombre']

class ClienteDetailView(LoginRequiredMixin, DetailView):
    model = Cliente
    template_name = 'inventario/cliente_detail.html'
    context_object_name = 'cliente'

class ClienteCreateView(LoginRequiredMixin, CreateView):
    model = Cliente
    form_class = ClienteForm
    template_name = 'inventario/cliente_form.html'
    success_url = reverse_lazy('inventario:cliente_list')

class ClienteUpdateView(LoginRequiredMixin, UpdateView):
    model = Cliente
    form_class = ClienteForm
    template_name = 'inventario/cliente_form.html'
    success_url = reverse_lazy('inventario:cliente_list')

class ClienteDeleteView(LoginRequiredMixin, DeleteView):
    model = Cliente
    template_name = 'inventario/cliente_confirm_delete.html'
    success_url = reverse_lazy('inventario:cliente_list')

# --- Vistas para Cotizaciones ---
DetalleCotizacionFormSet = inlineformset_factory(Cotizacion, DetalleCotizacion, form=DetalleCotizacionForm, extra=1, can_delete=True)

class CotizacionListView(LoginRequiredMixin, ListView):
    model = Cotizacion
    template_name = 'inventario/cotizacion_list.html'
    context_object_name = 'cotizaciones'
    ordering = ['-fecha_creacion']

class CotizacionDetailView(LoginRequiredMixin, DetailView):
    model = Cotizacion
    template_name = 'inventario/cotizacion_detail.html'
    context_object_name = 'cotizacion'

class CotizacionCreateView(LoginRequiredMixin, CreateView):
    model = Cotizacion
    form_class = CotizacionForm
    template_name = 'inventario/cotizacion_form.html'
    success_url = reverse_lazy('inventario:cotizacion_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.POST:
            context['detalle_formset'] = DetalleCotizacionFormSet(self.request.POST, instance=self.object)
        else:
            context['detalle_formset'] = DetalleCotizacionFormSet(instance=self.object)
        return context

    def form_valid(self, form):
        context = self.get_context_data()
        detalle_formset = context['detalle_formset']
        with transaction.atomic():
            self.object = form.save()
            if detalle_formset.is_valid():
                detalle_formset.instance = self.object
                detalles = detalle_formset.save(commit=False)
                total_cotizado = 0
                for detalle in detalles:
                    if detalle.producto:
                        detalle.nombre_producto_cotizado = detalle.producto.nombre
                        detalle.marca = detalle.producto.marca
                        detalle.modelo = detalle.producto.modelo
                        detalle.precio_unitario = detalle.producto.precio
                    detalle.subtotal = detalle.cantidad * detalle.precio_unitario
                    total_cotizado += detalle.subtotal
                    detalle.save()
                form.instance.total_cotizado = total_cotizado
                form.instance.save(update_fields=['total_cotizado'])
            else:
                return self.form_invalid(form)
        return super().form_valid(form)

class CotizacionUpdateView(LoginRequiredMixin, UpdateView):
    model = Cotizacion
    form_class = CotizacionForm
    template_name = 'inventario/cotizacion_form.html'
    success_url = reverse_lazy('inventario:cotizacion_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.POST:
            context['detalle_formset'] = DetalleCotizacionFormSet(self.request.POST, instance=self.object)
        else:
            context['detalle_formset'] = DetalleCotizacionFormSet(instance=self.object)
        return context

    def form_valid(self, form):
        context = self.get_context_data()
        detalle_formset = context['detalle_formset']
        with transaction.atomic():
            self.object = form.save()
            if detalle_formset.is_valid():
                detalle_formset.instance = self.object
                detalles = detalle_formset.save(commit=False)
                total_cotizado = 0
                for detalle in detalles:
                    if detalle.producto:
                        detalle.nombre_producto_cotizado = detalle.producto.nombre
                        detalle.marca = detalle.producto.marca
                        detalle.modelo = detalle.producto.modelo
                        detalle.precio_unitario = detalle.producto.precio
                    detalle.subtotal = detalle.cantidad * detalle.precio_unitario
                    total_cotizado += detalle.subtotal
                    detalle.save()
                form.instance.total_cotizado = total_cotizado
                form.instance.save(update_fields=['total_cotizado'])
            else:
                return self.form_invalid(form)
        return super().form_valid(form)

class CotizacionDeleteView(LoginRequiredMixin, DeleteView):
    model = Cotizacion
    template_name = 'inventario/cotizacion_confirm_delete.html'
    success_url = reverse_lazy('inventario:cotizacion_list')

def cotizacion_accept(request, pk):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    cotizacion = get_object_or_404(Cotizacion, pk=pk)
    if request.method == 'POST':
        if cotizacion.estado == 'abierta':
            with transaction.atomic():
                # Create OrdenSalida as a record of the accepted quote
                orden_salida = OrdenSalida.objects.create(
                    cliente=cotizacion.cliente.nombre,
                    cotizacion_origen=cotizacion,
                    total=cotizacion.total_cotizado
                )
                # Change cotizacion status to 'aceptada'
                cotizacion.estado = 'aceptada'
                cotizacion.save(update_fields=['estado'])

                # Redirect to the NotaDespacho creation form, pre-filling with this order
                return redirect('inventario:nota_despacho_create_from_order', orden_salida_id=orden_salida.pk)
        else:
            return HttpResponse("La cotización no está en estado 'abierta' para ser aceptada.", status=400)
    return HttpResponse("Método no permitido.", status=405)


class CotizacionesInterfaceView(LoginRequiredMixin, TemplateView):
    template_name = 'inventario/cotizaciones_interface.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['productos'] = Producto.objects.filter(cantidad__gt=0).order_by('nombre')
        context['clientes'] = Cliente.objects.all().order_by('nombre')
        return context

@csrf_exempt
def create_cotizacion_from_interface(request):
    if not request.user.is_authenticated:
        return JsonResponse({'status': 'error', 'message': 'No autenticado.'}, status=401)
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            cliente_id = data.get('cliente_id')
            items = data.get('items', [])

            if not cliente_id:
                return JsonResponse({'status': 'error', 'message': 'Debe seleccionar un cliente.'}, status=400)
            
            try:
                cliente = Cliente.objects.get(id=cliente_id)
            except Cliente.DoesNotExist:
                return JsonResponse({'status': 'error', 'message': 'Cliente no encontrado.'}, status=400)

            if not items:
                return JsonResponse({'status': 'error', 'message': 'No hay ítems en la cotización.'}, status=400)

            total_cotizacion = 0
            detalles_para_cotizacion = []

            with transaction.atomic():
                for item_data in items:
                    product_id = item_data.get('producto_id')
                    cantidad_solicitada = item_data.get('cantidad')
                    precio_unitario = item_data.get('precio_unitario')
                    nombre_producto_cotizado = item_data.get('nombre_producto_cotizado')
                    marca = item_data.get('marca')
                    modelo = item_data.get('modelo')

                    if not product_id or not cantidad_solicitada or precio_unitario is None:
                        raise ValueError("Datos de ítem incompletos para la cotización.")
                    
                    try:
                        producto = Producto.objects.get(id=product_id)
                    except Producto.DoesNotExist:
                        raise ValueError(f"Producto con ID {product_id} no encontrado.")

                    if cantidad_solicitada <= 0:
                        raise ValueError(f"Cantidad inválida para {producto.nombre}.")
                    
                    subtotal_item = float(precio_unitario) * cantidad_solicitada
                    total_cotizacion += subtotal_item
                    detalles_para_cotizacion.append({
                        'producto': producto,
                        'cantidad': cantidad_solicitada,
                        'precio_unitario': float(precio_unitario),
                        'subtotal': subtotal_item,
                        'nombre_producto_cotizado': nombre_producto_cotizado,
                        'marca': marca,
                        'modelo': modelo
                    })
                
                cotizacion = Cotizacion.objects.create(
                    cliente=cliente,
                    total_cotizado=total_cotizacion,
                    estado='abierta'
                )

                for item_data in detalles_para_cotizacion:
                    DetalleCotizacion.objects.create(
                        cotizacion=cotizacion,
                        producto=item_data['producto'],
                        cantidad=item_data['cantidad'],
                        precio_unitario=item_data['precio_unitario'],
                        subtotal=item_data['subtotal'],
                        nombre_producto_cotizado=item_data['nombre_producto_cotizado'],
                        marca=item_data['marca'],
                        modelo=item_data['modelo']
                    )

            return JsonResponse({'status': 'success', 'message': f'Cotización {cotizacion.numero_cotizacion} creada con éxito.', 'cotizacion_id': cotizacion.id})

        except ValueError as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': f'Error interno del servidor: {e}'}, status=500)
    
    return JsonResponse({'status': 'error', 'message': 'Método no permitido.'}, status=405)

@csrf_exempt
def get_cotizacion_details_api(request, pk):
    if not request.user.is_authenticated:
        return JsonResponse({'status': 'error', 'message': 'No autenticado.'}, status=401)
    """
    API endpoint to get details of a Cotizacion, including its DetalleCotizacion items.
    """
    try:
        cotizacion = Cotizacion.objects.get(pk=pk)
        detalles = cotizacion.detalles.all().values(
            'id', 'producto_id', 'nombre_producto_cotizado', 'marca', 'modelo', 'cantidad', 'cantidad_entregada', 'precio_unitario'
        )
        return JsonResponse({'status': 'success', 'detalles': list(detalles)})
    except Cotizacion.DoesNotExist:
        return JsonResponse({'status': 'error', 'message': 'Cotización no encontrada.'}, status=404)
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)

# --- Vistas para ActaRecepcion ---
# Define the formset for DetalleActaRecepcion
DetalleActaRecepcionFormSet = inlineformset_factory(
    ActaRecepcion,
    DetalleActaRecepcion,
    form=DetalleActaRecepcionForm,
    extra=0, # Changed to 0 as JS will manage adding rows
    can_delete=True
)

class ActaRecepcionListView(LoginRequiredMixin, ListView):
    model = ActaRecepcion
    template_name = 'inventario/acta_recepcion_list.html' # Nueva plantilla
    context_object_name = 'actas_recepcion'
    ordering = ['-fecha_recepcion']

class ActaRecepcionDetailView(LoginRequiredMixin, DetailView):
    model = ActaRecepcion
    template_name = 'inventario/acta_recepcion_detail.html' # Nueva plantilla
    context_object_name = 'acta_recepcion'

class ActaRecepcionCreateView(LoginRequiredMixin, CreateView):
    model = ActaRecepcion
    form_class = ActaRecepcionForm
    template_name = 'inventario/acta_recepcion_form.html' # Nueva plantilla
    success_url = reverse_lazy('inventario:acta_recepcion_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.POST:
            context['detalle_formset'] = DetalleActaRecepcionFormSet(self.request.POST, instance=self.object)
        else:
            context['detalle_formset'] = DetalleActaRecepcionFormSet(instance=self.object)
        
        context['all_products'] = list(Producto.objects.all().values('id', 'nombre', 'sku', 'marca', 'modelo', 'cantidad', 'unidad_medida__abreviatura'))
        
        # Serializar los datos iniciales del formset a JSON para JavaScript
        initial_detalles_data = []
        if self.object and self.object.pk: # Si estamos editando una Acta existente
            for detalle in self.object.detalles.all():
                initial_detalles_data.append({
                    'id': detalle.id,
                    'producto': detalle.producto.id if detalle.producto else None,
                    'nombre_producto_recepcion': detalle.nombre_producto_recepcion,
                    'marca': detalle.marca,
                    'modelo': detalle.modelo,
                    'cantidad': detalle.cantidad,
                })
        context['detalle_formset'].initial_json = json.dumps(initial_detalles_data)

        return context

    def form_valid(self, form):
        context = self.get_context_data()
        detalle_formset = context['detalle_formset']
        with transaction.atomic():
            self.object = form.save()
            if detalle_formset.is_valid():
                detalle_formset.instance = self.object
                detalle_formset.save()
            else:
                return self.form_invalid(form)
        return super().form_valid(form)

class ActaRecepcionUpdateView(LoginRequiredMixin, UpdateView):
    model = ActaRecepcion
    form_class = ActaRecepcionForm
    template_name = 'inventario/acta_recepcion_form.html' # Nueva plantilla
    success_url = reverse_lazy('inventario:acta_recepcion_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.POST:
            context['detalle_formset'] = DetalleActaRecepcionFormSet(self.request.POST, instance=self.object)
        else:
            context['detalle_formset'] = DetalleActaRecepcionFormSet(instance=self.object)
        
        context['all_products'] = list(Producto.objects.all().values('id', 'nombre', 'sku', 'marca', 'modelo', 'cantidad', 'unidad_medida__abreviatura'))
        
        # Serializar los datos iniciales del formset a JSON para JavaScript
        initial_detalles_data = []
        if self.object and self.object.pk: # Si estamos editando una Acta existente
            for detalle in self.object.detalles.all():
                initial_detalles_data.append({
                    'id': detalle.id,
                    'producto': detalle.producto.id if detalle.producto else None,
                    'nombre_producto_recepcion': detalle.nombre_producto_recepcion,
                    'marca': detalle.marca,
                    'modelo': detalle.modelo,
                    'cantidad': detalle.cantidad,
                })
        context['detalle_formset'].initial_json = json.dumps(initial_detalles_data)

        return context

    def form_valid(self, form):
        context = self.get_context_data()
        detalle_formset = context['detalle_formset']
        with transaction.atomic():
            self.object = form.save()
            if detalle_formset.is_valid():
                detalle_formset.instance = self.object
                detalle_formset.save()
            else:
                return self.form_invalid(form)
        return super().form_valid(form)

class ActaRecepcionDeleteView(LoginRequiredMixin, DeleteView):
    model = ActaRecepcion
    template_name = 'inventario/acta_recepcion_confirm_delete.html' # Nueva plantilla
    success_url = reverse_lazy('inventario:acta_recepcion_list')

def export_acta_recepcion_pdf(request, pk):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    """
    Genera un Acta de Recepción en formato PDF utilizando xhtml2pdf.
    """
    try:
        acta_recepcion = get_object_or_404(ActaRecepcion, pk=pk)
    except Http404:
        return HttpResponse("Acta de Recepción no encontrada.", status=404)

    detalles = acta_recepcion.detalles.all()
    # Calcular cuántas filas vacías necesitamos para llegar a 5 ítems (como en la plantilla)
    num_empty_rows = max(0, 5 - len(detalles))
    empty_rows_range = range(num_empty_rows)

    # --- Obtener la ruta absoluta de la imagen ---
    image_static_path = 'img/encabezado_oficial.jpg'
    absolute_image_path = None
    try:
        absolute_image_path = staticfiles_storage.path(image_static_path)
        print(f"Ruta absoluta de la imagen para PDF (Acta): {absolute_image_path}") # Debugging
        if not os.path.exists(absolute_image_path):
            print(f"ADVERTENCIA (Acta): La imagen no existe en la ruta absoluta: {absolute_image_path}")
            absolute_image_path = None
    except NotImplementedError:
        print("ADVERTENCIA (Acta): staticfiles_storage.path() no soportado. Intentando con STATIC_ROOT.")
        if hasattr(settings, 'STATIC_ROOT') and settings.STATIC_ROOT:
            absolute_image_path = os.path.join(settings.STATIC_ROOT, image_static_path)
            if not os.path.exists(absolute_image_path):
                print(f"ADVERTENCIA (Acta): La imagen no existe en STATIC_ROOT: {absolute_image_path}")
                absolute_image_path = None
        else:
            print("ADVERTENCIA (Acta): STATIC_ROOT no está configurado, no se puede obtener la ruta absoluta de la imagen.")
            absolute_image_path = None
    except Exception as e:
        print(f"Error al obtener la ruta absoluta de la imagen (Acta): {e}")
        absolute_image_path = None

    context = {
        'acta_recepcion': acta_recepcion,
        'detalles': detalles,
        'total_cantidad': sum(detalle.cantidad for detalle in detalles),
        'current_date': timezone.now().strftime("%d/%m/%Y"),
        'empty_rows_range': empty_rows_range,
        'header_image_path': absolute_image_path,
    }

    template_path = 'inventario/pdf/acta_recepcion_pdf.html'
    html = render_to_string(template_path, context)

    pisa_status = pisa.CreatePDF(
        html,
        dest=response)

    if pisa_status.err:
        return HttpResponse('Error al generar el PDF: %s' % pisa_status.err, status=500)
    return response

class OrdenSalidaDetailView(LoginRequiredMixin, DetailView):
    model = OrdenSalida
    template_name = 'inventario/orden_salida_detail.html'
    context_object_name = 'orden_salida'

# --- Vistas para Clientes ---
class ClienteListView(LoginRequiredMixin, ListView):
    model = Cliente
    template_name = 'inventario/cliente_list.html'
    context_object_name = 'clientes'
    ordering = ['nombre']

class ClienteDetailView(LoginRequiredMixin, DetailView):
    model = Cliente
    template_name = 'inventario/cliente_detail.html'
    context_object_name = 'cliente'

class ClienteCreateView(LoginRequiredMixin, CreateView):
    model = Cliente
    form_class = ClienteForm
    template_name = 'inventario/cliente_form.html'
    success_url = reverse_lazy('inventario:cliente_list')

class ClienteUpdateView(LoginRequiredMixin, UpdateView):
    model = Cliente
    form_class = ClienteForm
    template_name = 'inventario/cliente_form.html'
    success_url = reverse_lazy('inventario:cliente_list')

class ClienteDeleteView(LoginRequiredMixin, DeleteView):
    model = Cliente
    template_name = 'inventario/cliente_confirm_delete.html'
    success_url = reverse_lazy('inventario:cliente_list')

# --- Vistas para Cotizaciones ---
DetalleCotizacionFormSet = inlineformset_factory(Cotizacion, DetalleCotizacion, form=DetalleCotizacionForm, extra=1, can_delete=True)

class CotizacionListView(LoginRequiredMixin, ListView):
    model = Cotizacion
    template_name = 'inventario/cotizacion_list.html'
    context_object_name = 'cotizaciones'
    ordering = ['-fecha_creacion']

class CotizacionDetailView(LoginRequiredMixin, DetailView):
    model = Cotizacion
    template_name = 'inventario/cotizacion_detail.html'
    context_object_name = 'cotizacion'

class CotizacionCreateView(LoginRequiredMixin, CreateView):
    model = Cotizacion
    form_class = CotizacionForm
    template_name = 'inventario/cotizacion_form.html'
    success_url = reverse_lazy('inventario:cotizacion_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.POST:
            context['detalle_formset'] = DetalleCotizacionFormSet(self.request.POST, instance=self.object)
        else:
            context['detalle_formset'] = DetalleCotizacionFormSet(instance=self.object)
        return context

    def form_valid(self, form):
        context = self.get_context_data()
        detalle_formset = context['detalle_formset']
        with transaction.atomic():
            self.object = form.save()
            if detalle_formset.is_valid():
                detalle_formset.instance = self.object
                detalles = detalle_formset.save(commit=False)
                total_cotizado = 0
                for detalle in detalles:
                    if detalle.producto:
                        detalle.nombre_producto_cotizado = detalle.producto.nombre
                        detalle.marca = detalle.producto.marca
                        detalle.modelo = detalle.producto.modelo
                        detalle.precio_unitario = detalle.producto.precio
                    detalle.subtotal = detalle.cantidad * detalle.precio_unitario
                    total_cotizado += detalle.subtotal
                    detalle.save()
                form.instance.total_cotizado = total_cotizado
                form.instance.save(update_fields=['total_cotizado'])
            else:
                return self.form_invalid(form)
        return super().form_valid(form)

class CotizacionUpdateView(LoginRequiredMixin, UpdateView):
    model = Cotizacion
    form_class = CotizacionForm
    template_name = 'inventario/cotizacion_form.html'
    success_url = reverse_lazy('inventario:cotizacion_list')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.POST:
            context['detalle_formset'] = DetalleCotizacionFormSet(self.request.POST, instance=self.object)
        else:
            context['detalle_formset'] = DetalleCotizacionFormSet(instance=self.object)
        return context

    def form_valid(self, form):
        context = self.get_context_data()
        detalle_formset = context['detalle_formset']
        with transaction.atomic():
            self.object = form.save()
            if detalle_formset.is_valid():
                detalle_formset.instance = self.object
                detalles = detalle_formset.save(commit=False)
                total_cotizado = 0
                for detalle in detalles:
                    if detalle.producto:
                        detalle.nombre_producto_cotizado = detalle.producto.nombre
                        detalle.marca = detalle.producto.marca
                        detalle.modelo = detalle.producto.modelo
                        detalle.precio_unitario = detalle.producto.precio
                    detalle.subtotal = detalle.cantidad * detalle.precio_unitario
                    total_cotizado += detalle.subtotal
                    detalle.save()
                form.instance.total_cotizado = total_cotizado
                form.instance.save(update_fields=['total_cotizado'])
            else:
                return self.form_invalid(form)
        return super().form_valid(form)

class CotizacionDeleteView(LoginRequiredMixin, DeleteView):
    model = Cotizacion
    template_name = 'inventario/cotizacion_confirm_delete.html'
    success_url = reverse_lazy('inventario:cotizacion_list')

def cotizacion_accept(request, pk):
    if not request.user.is_authenticated:
        return redirect('inventario:login') # Redirigir si no está autenticado
    cotizacion = get_object_or_404(Cotizacion, pk=pk)
    if request.method == 'POST':
        if cotizacion.estado == 'abierta':
            with transaction.atomic():
                # Create OrdenSalida as a record of the accepted quote
                orden_salida = OrdenSalida.objects.create(
                    cliente=cotizacion.cliente.nombre,
                    cotizacion_origen=cotizacion,
                    total=cotizacion.total_cotizado
                )
                # Change cotizacion status to 'aceptada'
                cotizacion.estado = 'aceptada'
                cotizacion.save(update_fields=['estado'])

                # Redirect to the NotaDespacho creation form, pre-filling with this order
                return redirect('inventario:nota_despacho_create_from_order', orden_salida_id=orden_salida.pk)
        else:
            return HttpResponse("La cotización no está en estado 'abierta' para ser aceptada.", status=400)
    return HttpResponse("Método no permitido.", status=405)


class CotizacionesInterfaceView(LoginRequiredMixin, TemplateView):
    template_name = 'inventario/cotizaciones_interface.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['productos'] = Producto.objects.filter(cantidad__gt=0).order_by('nombre')
        context['clientes'] = Cliente.objects.all().order_by('nombre')
        return context

@csrf_exempt
def create_cotizacion_from_interface(request):
    if not request.user.is_authenticated:
        return JsonResponse({'status': 'error', 'message': 'No autenticado.'}, status=401)
    if request.method == 'POST':
        try:
            data = json.loads(request.body)
            cliente_id = data.get('cliente_id')
            items = data.get('items', [])

            if not cliente_id:
                return JsonResponse({'status': 'error', 'message': 'Debe seleccionar un cliente.'}, status=400)
            
            try:
                cliente = Cliente.objects.get(id=cliente_id)
            except Cliente.DoesNotExist:
                return JsonResponse({'status': 'error', 'message': 'Cliente no encontrado.'}, status=400)

            if not items:
                return JsonResponse({'status': 'error', 'message': 'No hay ítems en la cotización.'}, status=400)

            total_cotizacion = 0
            detalles_para_cotizacion = []

            with transaction.atomic():
                for item_data in items:
                    product_id = item_data.get('producto_id')
                    cantidad_solicitada = item_data.get('cantidad')
                    precio_unitario = item_data.get('precio_unitario')
                    nombre_producto_cotizado = item_data.get('nombre_producto_cotizado')
                    marca = item_data.get('marca')
                    modelo = item_data.get('modelo')

                    if not product_id or not cantidad_solicitada or precio_unitario is None:
                        raise ValueError("Datos de ítem incompletos para la cotización.")
                    
                    try:
                        producto = Producto.objects.get(id=product_id)
                    except Producto.DoesNotExist:
                        raise ValueError(f"Producto con ID {product_id} no encontrado.")

                    if cantidad_solicitada <= 0:
                        raise ValueError(f"Cantidad inválida para {producto.nombre}.")
                    
                    subtotal_item = float(precio_unitario) * cantidad_solicitada
                    total_cotizacion += subtotal_item
                    detalles_para_cotizacion.append({
                        'producto': producto,
                        'cantidad': cantidad_solicitada,
                        'precio_unitario': float(precio_unitario),
                        'subtotal': subtotal_item,
                        'nombre_producto_cotizado': nombre_producto_cotizado,
                        'marca': marca,
                        'modelo': modelo
                    })
                
                cotizacion = Cotizacion.objects.create(
                    cliente=cliente,
                    total_cotizado=total_cotizacion,
                    estado='abierta'
                )

                for item_data in detalles_para_cotizacion:
                    DetalleCotizacion.objects.create(
                        cotizacion=cotizacion,
                        producto=item_data['producto'],
                        cantidad=item_data['cantidad'],
                        precio_unitario=item_data['precio_unitario'],
                        subtotal=item_data['subtotal'],
                        nombre_producto_cotizado=item_data['nombre_producto_cotizado'],
                        marca=item_data['marca'],
                        modelo=item_data['modelo']
                    )

            return JsonResponse({'status': 'success', 'message': f'Cotización {cotizacion.numero_cotizacion} creada con éxito.', 'cotizacion_id': cotizacion.id})

        except ValueError as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=400)
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': f'Error interno del servidor: {e}'}, status=500)
    
    return JsonResponse({'status': 'error', 'message': 'Método no permitido.'}, status=405)

@csrf_exempt
def get_cotizacion_details_api(request, pk):
    if not request.user.is_authenticated:
        return JsonResponse({'status': 'error', 'message': 'No autenticado.'}, status=401)
    """
    API endpoint to get details of a Cotizacion, including its DetalleCotizacion items.
    """
    try:
        cotizacion = Cotizacion.objects.get(pk=pk)
        detalles = cotizacion.detalles.all().values(
            'id', 'producto_id', 'nombre_producto_cotizado', 'marca', 'modelo', 'cantidad', 'cantidad_entregada', 'precio_unitario'
        )
        return JsonResponse({'status': 'success', 'detalles': list(detalles)})
    except Cotizacion.DoesNotExist:
        return JsonResponse({'status': 'error', 'message': 'Cotización no encontrada.'}, status=404)
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)

# --- Vistas para Pagos a Proveedores ---
class PagoProveedorListView(LoginRequiredMixin, ListView):
    model = PagoProveedor
    template_name = 'inventario/pago_proveedor_list.html'
    context_object_name = 'pagos'
    ordering = ['-fecha_pago']

class PagoProveedorDetailView(LoginRequiredMixin, DetailView):
    model = PagoProveedor
    template_name = 'inventario/pago_proveedor_detail.html'
    context_object_name = 'pago'

class PagoProveedorCreateView(LoginRequiredMixin, CreateView):
    model = PagoProveedor
    form_class = PagoProveedorForm
    template_name = 'inventario/pago_proveedor_form.html'
    success_url = reverse_lazy('inventario:pago_proveedor_list')

class PagoProveedorUpdateView(LoginRequiredMixin, UpdateView):
    model = PagoProveedor
    form_class = PagoProveedorForm
    template_name = 'inventario/pago_proveedor_form.html'
    success_url = reverse_lazy('inventario:pago_proveedor_list')

class PagoProveedorDeleteView(LoginRequiredMixin, DeleteView):
    model = PagoProveedor
    template_name = 'inventario/pago_proveedor_confirm_delete.html'
    success_url = reverse_lazy('inventario:pago_proveedor_list')

# --- Vista para el Panel de Gestión de Proveedores ---
class ProveedorManagementView(LoginRequiredMixin, TemplateView):
    template_name = 'inventario/proveedor_management_panel.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        
        proveedores = Proveedor.objects.all().order_by('nombre')
        proveedores_data = []

        for proveedor in proveedores:
            # Total pagado al proveedor
            total_pagado = PagoProveedor.objects.filter(proveedor=proveedor).aggregate(Sum('monto'))['monto__sum'] or 0

            # Despachos pendientes por recibir (Actas de Recepción con estado 'pendiente' o 'parcial')
            actas_pendientes = ActaRecepcion.objects.filter(
                proveedor=proveedor
            ).filter(
                Q(estado='pendiente') | Q(estado='parcial')
            ).order_by('-fecha_recepcion')

            # Despachos entregados (Actas de Recepción con estado 'completa')
            actas_entregadas = ActaRecepcion.objects.filter(
                proveedor=proveedor,
                estado='completa'
            ).order_by('-fecha_recepcion')

            # Valor de la mercancía recibida de este proveedor
            # Sumar el precio de los productos en las Actas de Recepción de este proveedor
            # Multiplicar la cantidad recibida por el precio actual del producto
            # Esto puede ser un precio de referencia o el precio al momento de la recepción si se guardó
            # Para simplificar, usaremos el precio actual del producto.
            total_mercancia_recibida_valor = 0
            actas_recibidas = ActaRecepcion.objects.filter(proveedor=proveedor).prefetch_related('detalles__producto')
            for acta in actas_recibidas:
                for detalle in acta.detalles.all():
                    if detalle.producto:
                        total_mercancia_recibida_valor += detalle.cantidad * detalle.producto.precio # Usar el precio actual del producto

            proveedores_data.append({
                'proveedor': proveedor,
                'total_pagado': total_pagado,
                'actas_pendientes': actas_pendientes,
                'actas_entregadas': actas_entregadas,
                'total_mercancia_recibida_valor': total_mercancia_recibida_valor,
            })
        
        context['proveedores_data'] = proveedores_data
        return context
    
def movimiento_crear(request):
    productos = Producto.objects.all().select_related('unidad_medida', 'almacen').order_by('nombre')
    tipos_movimiento = TipoMovimiento.objects.all().order_by('nombre')
    almacenes = Almacen.objects.all().order_by('nombre')

    # Convertir productos a un formato JSON seguro para JS
    # Ya no necesitas all_products_data JSON.dumps aquí si lo pasas al template
    # y lo construyes en JS como en cotizaciones_interface.html
    # La clave es pasar el queryset 'productos' directamente.

    context = {
        'form': MovimientoForm(), # Si usas un formulario, si no, puedes eliminarlo
        'productos': productos,
        'tipos_movimiento': tipos_movimiento,
        'almacenes': almacenes,
        'selected_products_initial_json': '[]' # Para iniciar sin productos seleccionados si es una creación
    }
    return render(request, 'inventario/movimiento_crear.html', context)



@csrf_exempt # Solo para desarrollo, usar el token CSRF en producción
def crear_movimiento_from_interface(request):
    try:
        data = json.loads(request.body)
        tipo_movimiento_id = data.get('tipo_movimiento_id')
        almacen_origen_id = data.get('almacen_origen_id')
        almacen_destino_id = data.get('almacen_destino_id')
        descripcion = data.get('descripcion', '')
        items = data.get('items', [])

        if not tipo_movimiento_id or not almacen_origen_id or not items:
            return JsonResponse({'status': 'error', 'message': 'Faltan datos requeridos (tipo de movimiento, almacén de origen o ítems).'}, status=400)

        tipo_movimiento = get_object_or_404(TipoMovimiento, id=tipo_movimiento_id)
        almacen_origen = get_object_or_404(Almacen, id=almacen_origen_id)
        almacen_destino = get_object_or_404(Almacen, id=almacen_destino_id) if almacen_destino_id else None

        with transaction.atomic():
            movimiento = Movimiento.objects.create(
                tipo_movimiento=tipo_movimiento,
                almacen_origen=almacen_origen,
                almacen_destino=almacen_destino,
                descripcion=descripcion,
                responsable=request.user # Asigna el usuario logueado como responsable
            )

            for item_data in items:
                product_id = item_data.get('producto_id')
                cantidad = item_data.get('cantidad')

                if not product_id or not cantidad or cantidad <= 0:
                    raise ValueError("Datos de ítem inválidos.")

                producto = get_object_or_404(Producto, id=product_id)

                # Lógica de stock
                if tipo_movimiento.es_entrada:
                    producto.cantidad += cantidad
                elif tipo_movimiento.es_salida:
                    if producto.cantidad < cantidad:
                        raise ValueError(f"No hay suficiente stock para {producto.nombre}. Disponible: {producto.cantidad}, Solicitado: {cantidad}")
                    producto.cantidad -= cantidad
                elif tipo_movimiento.es_traslado:
                    if producto.cantidad < cantidad:
                        raise ValueError(f"No hay suficiente stock en origen para {producto.nombre}. Disponible: {producto.cantidad}, Solicitado: {cantidad}")
                    # En traslados, restamos del origen y sumamos al destino
                    producto.cantidad -= cantidad
                    if almacen_destino:
                        # Buscar o crear el producto en el almacén destino
                        # Esto es una simplificación. Un manejo más robusto podría requerir:
                        # 1. Chequear si el producto existe ya en el almacén destino con ese ID.
                        # 2. Si no, decidir si crear una nueva instancia de Producto
                        #    o manejar el stock de forma global.
                        # Por ahora, simplemente actualizamos la cantidad del producto original.
                        # Si Producto se relaciona con Almacen, necesitarías un objeto ProductoAlmacen.
                        # Para este ejemplo, asumimos que 'cantidad' en Producto es global o que el traslado
                        # implica mover la misma instancia de producto entre almacenes.
                        # SI tienes un modelo ProductoAlmacen (MUY RECOMENDADO para múltiples almacenes):
                        # producto_almacen_destino, created = ProductoAlmacen.objects.get_or_create(
                        #     producto=producto, almacen=almacen_destino, defaults={'cantidad': 0}
                        # )
                        # producto_almacen_destino.cantidad += cantidad
                        # producto_almacen_destino.save()
                        pass # La lógica de movimiento de stock entre almacenes complejos va aquí
                    else:
                        raise ValueError("Almacén destino es requerido para movimientos de tipo 'Traslado'.")

                producto.save()

                DetalleMovimiento.objects.create(
                    movimiento=movimiento,
                    producto=producto,
                    cantidad=cantidad,
                    precio_unitario_al_momento=producto.precio # Guarda el precio del producto al momento del movimiento
                )

            return JsonResponse({'status': 'success', 'message': 'Movimiento creado exitosamente.', 'movimiento_id': movimiento.id})

    except ValueError as e:
        transaction.set_rollback(True)
        return JsonResponse({'status': 'error', 'message': str(e)}, status=400)
    except Exception as e:
        transaction.set_rollback(True)
        print(f"Error inesperado al crear movimiento: {e}")
        return JsonResponse({'status': 'error', 'message': 'Ocurrió un error inesperado al crear el movimiento.'}, status=500)