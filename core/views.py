import io
import uuid
import qrcode
import openpyxl
from itertools import groupby  # Importación clave para agrupar en el backend
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Sum, Count, F, DecimalField
from django.db.models.functions import Coalesce
from django.template.loader import get_template

# Renderizado de PDF
from xhtml2pdf import pisa
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

# Modelos y formularios del núcleo
from .models import (
    Producto, Proveedor, Cliente, Cotizacion,
    ItemCotizacion, ActaRecepcion, NotaDespacho,
    ItemDespacho, PagoProveedor, Almacen
)
from .forms import (
    ProductoForm,
    CotizacionForm,
    ItemCotizacionFormSet,
    ActaRecepcionForm,
    MovimientoRecepcionFormSet,
    NotaDespachoForm,
    ItemDespachoFormSet,
    NotaDevolucionForm,
    NotaDevolucionItemFormSet,
    ClienteForm,
    ProveedorForm,
    AlmacenForm,
)

# ==============================================================================
# PROYECTO: EMVEPRO C.A.
# GERENCIA DE TECNOLOGÍA DE LA INFORMACIÓN Y COMUNICACIÓN
# ==============================================================================

@login_required
def dashboard_principal(request):
    """Dashboard Corporativo con métricas precisas en USD/Bs."""
    total_productos = Producto.objects.filter(activo=True).count()
    productos_bajo_stock = Producto.objects.filter(activo=True, stock_actual__lte=0.00).count()

    total_cotizaciones = Cotizacion.objects.count()
    total_actas_recepcion = ActaRecepcion.objects.count()
    total_notas_despacho = NotaDespacho.objects.count()

    totales_items = ItemCotizacion.objects.aggregate(
        total_proyectado_usd=Sum('total_item_usd')
    )
    monto_proyectado_usd = totales_items['total_proyectado_usd'] or 0.00

    cotizaciones_pendientes = Cotizacion.objects.filter(estatus='PENDIENTE').count()
    cotizaciones_dotadas = Cotizacion.objects.filter(estatus='DOTADO').count()

    financiero_global = Cotizacion.objects.exclude(estatus='ANULADO').aggregate(
        acumulado_bs=Sum('total_bs'),
        acumulado_usd=Sum('total_usd')
    )
    total_general_bs = financiero_global['acumulado_bs'] or 0.00
    total_general_usd = financiero_global['acumulado_usd'] or 0.00

    cotizaciones_recientes = Cotizacion.objects.select_related('cliente').order_by('-fecha_emision')[:5]
    despachos_recientes = NotaDespacho.objects.select_related('cotizacion', 'responsable').order_by('-fecha_despacho')[:5]

    return render(request, 'core/dashboard.html', {
        'total_productos': total_productos,
        'productos_bajo_stock': productos_bajo_stock,
        'total_cotizaciones': total_cotizaciones,
        'total_actas_recepcion': total_actas_recepcion,
        'total_notas_despacho': total_notas_despacho,
        'cotizaciones_pendientes': cotizaciones_pendientes,
        'cotizaciones_dotadas': cotizaciones_dotadas,
        'monto_proyectado_usd': monto_proyectado_usd,
        'total_general_bs': total_general_bs,
        'total_general_usd': total_general_usd,
        'cotizaciones_recientes': cotizaciones_recientes,
        'despachos_recientes': despachos_recientes,
    })

@login_required
def estatus_proyectos(request):
    """Porcentaje de dotación de cada proyecto activo."""
    cotizaciones = Cotizacion.objects.exclude(estatus='ANULADO').prefetch_related('items__despachos')
    proyectos_data = []

    for cot in cotizaciones:
        total_solicitado = sum(item.cantidad_solicitada for item in cot.items.all())
        total_despachado = sum(
            sum(despacho.cantidad_despachada for despacho in item.despachos.all())
            for item in cot.items.all()
        )
        porcentaje = (total_despachado / total_solicitado * 100) if total_solicitado > 0 else 0

        proyectos_data.append({
            'id': cot.id,
            'numero_rastreo': cot.numero_rastreo,
            'cliente': cot.cliente.nombre,
            'estatus': cot.get_estatus_display(),
            'porcentaje_dotacion': round(porcentaje, 2),
        })

    return render(request, 'core/estatus_proyectos.html', {'proyectos': proyectos_data})

@login_required
def finanzas_proveedores(request):
    """Cuentas por pagar a proveedores."""
    proveedores = Proveedor.objects.all()
    datos_financieros = []

    for prov in proveedores:
        facturado = prov.actas_recepcion.aggregate(
            total=Coalesce(Sum('monto_total_factura', output_field=DecimalField()), 0)
        )['total']
        pagado = prov.pagos.aggregate(
            total=Coalesce(Sum('monto_pagado', output_field=DecimalField()), 0)
        )['total']
        deuda = facturado - pagado

        if facturado > 0 or pagado > 0:
            datos_financieros.append({
                'proveedor': prov.nombre,
                'total_facturado': facturado,
                'total_pagado': pagado,
                'deuda_actual': deuda,
            })

    return render(request, 'core/finanzas_proveedores.html', {'datos_financieros': datos_financieros})

@login_required
def lista_productos(request):
    productos = Producto.objects.select_related('almacen').all().order_by('nombre')
    return render(request, 'core/producto_list.html', {'productos': productos})

@login_required
def crear_producto(request):
    if request.method == 'POST':
        form = ProductoForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('core:lista_productos')
    else:
        form = ProductoForm()
    return render(request, 'core/producto_form.html', {'form': form})

# ==============================================================================
# MÓDULO COMERCIAL: COTIZACIONES
# ==============================================================================

@login_required
def lista_cotizaciones(request):
    """Listado general de cotizaciones con accesos a PDF."""
    cotizaciones = Cotizacion.objects.select_related('cliente').all().order_by('-fecha_emision')
    return render(request, 'core/cotizacion_list.html', {'cotizaciones': cotizaciones})

@login_required
def crear_cotizacion(request):
    """Crea una cotización nueva usando el POS."""
    if request.method == 'POST':
        form = CotizacionForm(request.POST)
        formset = ItemCotizacionFormSet(request.POST)
        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                cotizacion = form.save()
                formset.instance = cotizacion
                formset.save()
            return redirect('core:lista_cotizaciones')
    else:
        form = CotizacionForm()
        formset = ItemCotizacionFormSet()
    return render(request, 'core/cotizacion_form.html', {'form': form, 'formset': formset})

@login_required
def editar_cotizacion(request, cotizacion_id):
    """Edita una cotización existente reutilizando el formset dinámico del POS."""
    cotizacion = get_object_or_404(Cotizacion, id=cotizacion_id)
    
    if request.method == 'POST':
        form = CotizacionForm(request.POST, instance=cotizacion)
        formset = ItemCotizacionFormSet(request.POST, instance=cotizacion)
        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                instancia_cot = form.save(commit=False)
                instancia_cot.save()
                formset.save()
                
                # Forzar recalculo por si cambiaron tasas o IVA
                instancia_cot.calcular_totales_reales()
                instancia_cot.save()
                
            return redirect('core:lista_cotizaciones')
    else:
        form = CotizacionForm(instance=cotizacion)
        formset = ItemCotizacionFormSet(instance=cotizacion)
        
    return render(request, 'core/cotizacion_form.html', {
        'form': form,
        'formset': formset,
        'cotizacion': cotizacion
    })

# ==============================================================================
# MÓDULO LOGÍSTICO: RECEPCIÓN, DESPACHO Y DEVOLUCIÓN
# ==============================================================================

@login_required
def crear_acta_recepcion(request):
    if request.method == 'POST':
        form = ActaRecepcionForm(request.POST)
        formset = MovimientoRecepcionFormSet(request.POST)
        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                acta = form.save()
                movimientos = formset.save(commit=False)
                for mov in movimientos:
                    mov.acta_recepcion = acta
                    mov.almacen = acta.almacen
                    mov.responsable = acta.responsable
                    mov.tipo = 'ENTRADA'
                    mov.save()
                    mov.producto.stock_actual += mov.cantidad
                    mov.producto.save()
            return redirect('core:dashboard')
    else:
        form = ActaRecepcionForm()
        formset = MovimientoRecepcionFormSet()
    return render(request, 'core/recepcion_form.html', {'form': form, 'formset': formset})

@login_required
def crear_nota_despacho(request):
    if request.method == 'POST':
        form = NotaDespachoForm(request.POST)
        formset = ItemDespachoFormSet(request.POST)
        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                nota_despacho = form.save(commit=False)
                nota_despacho.responsable = request.user
                nota_despacho.save()
                formset.instance = nota_despacho
                formset.save()
            return redirect('core:dashboard')
    else:
        form = NotaDespachoForm()
        formset = ItemDespachoFormSet()
    return render(request, 'core/despacho_form.html', {'form': form, 'formset': formset})

@login_required
def crear_nota_devolucion(request):
    if request.method == 'POST':
        form = NotaDevolucionForm(request.POST)
        formset = NotaDevolucionItemFormSet(request.POST)
        if form.is_valid() and formset.is_valid():
            with transaction.atomic():
                devolucion = form.save(commit=False)
                devolucion.responsable = request.user
                devolucion.save()
                formset.instance = devolucion
                formset.save()
            return redirect('core:estatus_proyectos')
    else:
        form = NotaDevolucionForm()
        formset = NotaDevolucionItemFormSet()
    return render(request, 'core/crear_nota_devolucion.html', {'form': form, 'formset': formset})

# ==============================================================================
# GENERADORES DE REPORTES PDF
# ==============================================================================

@login_required
def generar_pdf_cotizacion(request, cotizacion_id):
    """Estilo 1: PDF de Cotización ordenado para {% regroup %} en HTML."""
    cotizacion = get_object_or_404(Cotizacion, id=cotizacion_id)
    
    # Extraemos los ítems ordenados en una variable independiente
    items_planos = cotizacion.items.all().order_by('seccion_departamento')
    
    # Pasamos un diccionario de contexto explícito y robusto
    context = {
        'cotizacion': cotizacion,
        'items_ordenados': items_planos
    }
    
    template = get_template('core/cotizacion_pdf.html')
    html = template.render(context) # <--- Enviamos el diccionario completo
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="Cotizacion_EMVEPRO_{cotizacion.numero_rastreo}.pdf"'
    
    pisa_status = pisa.CreatePDF(html, dest=response)
    if pisa_status.err:
        return HttpResponse('Error generando el PDF Estándar', status=500)
    return response

@login_required
def generar_pdf_cotizacion_2(request, cotizacion_id):
    """Estilo 2: PDF con Subtotales agrupados matemáticamente en el backend."""
    cotizacion = get_object_or_404(Cotizacion, id=cotizacion_id)
    
    items = cotizacion.items.all().order_by('seccion_departamento')
    
    items_agrupados = []
    for seccion, grupo in groupby(items, key=lambda x: x.seccion_departamento):
        lista_items = list(grupo)
        subtotal_seccion = sum(item.total_item_usd for item in lista_items)
        
        items_agrupados.append({
            'seccion': seccion,
            'items': lista_items,
            'subtotal_seccion': subtotal_seccion
        })
    
    context = {
        'cotizacion': cotizacion,
        'items_agrupados': items_agrupados
    }
    
    template = get_template('core/cotizacion_2_pdf.html')
    html = template.render(context)
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="Cotizacion_Subtotales_{cotizacion.numero_rastreo}.pdf"'
    
    pisa_status = pisa.CreatePDF(html, dest=response)
    if pisa_status.err:
        return HttpResponse('Error generando el PDF de Subtotales', status=500)
    return response

@login_required
def generar_pdf_despacho(request, despacho_id):
    """PDF oficial de Nota de Despacho con QR dinámico."""
    despacho = get_object_or_404(NotaDespacho, id=despacho_id)
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="Nota_Despacho_{despacho.numero_guia}.pdf"'

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    story = []

    styles = getSampleStyleSheet()
    normal_style = ParagraphStyle('DocNormal', parent=styles['Normal'], fontName='Helvetica', fontSize=9, leading=12)
    bold_style = ParagraphStyle('DocBold', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9, leading=12)
    header_table_style = ParagraphStyle('TableHeader', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=9,
                                        textColor=colors.white, alignment=1)

    url_validacion = request.build_absolute_uri(f"/core/validar-despacho/{despacho.id}/")
    qr = qrcode.QRCode(version=1, box_size=3, border=1)
    qr.add_data(url_validacion)
    qr.make(fit=True)
    img_qr = qr.make_image(fill_color="black", back_color="white")
    qr_buffer = io.BytesIO()
    img_qr.save(qr_buffer, format='PNG')
    qr_buffer.seek(0)
    reportlab_qr = Image(qr_buffer, width=65, height=65)

    texto_institucion = """<b>Gobierno Bolivariano de Venezuela</b><br/>
    Ministerio del Poder Popular de Industrias y Producción Nacional<br/>
    Viceministerio de Industrias Intermedias y Ligeras<br/>
    <b>EMVEPRO C.A. / ALMACÉN FUERTE TIUNA</b>
    """
    info_guia = f"<b>NOTA DE DESPACHO</b><br/><font color='red'><b>Nro. {despacho.numero_guia}</b></font><br/>Fecha: {despacho.fecha_despacho.strftime('%d/%m/%Y %H:%M')}"

    data_top = [
        [Paragraph(texto_institucion, normal_style), reportlab_qr, Paragraph(info_guia, normal_style)]
    ]
    table_top = Table(data_top, colWidths=[260, 80, 210])
    table_top.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (2, 0), (2, 0), 'RIGHT'),
    ]))
    story.append(table_top)
    story.append(Spacer(1, 15))

    data_beneficiario = [
        [Paragraph("<b>NOMBRE DEL BENEFICIARIO:</b>", bold_style),
         Paragraph(despacho.nombre_beneficiario or "N/P", normal_style)],
        [Paragraph("<b>PROVEEDOR:</b>", bold_style),
         Paragraph(despacho.proveedor_origen or "EMVEPRO C.A.", normal_style)],
        [Paragraph("<b>N° ORDEN ASOCIADA / PROYECTO:</b>", bold_style),
         Paragraph(despacho.nro_orden_asociada or "N/P", normal_style)],
        [Paragraph("<b>NOTA DE ENTREGA ASOCIADA:</b>", bold_style),
         Paragraph(despacho.cotizacion.numero_rastreo, normal_style)]
    ]
    table_ben = Table(data_beneficiario, colWidths=[180, 370])
    table_ben.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F3F4F6')),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(table_ben)
    story.append(Spacer(1, 15))

    tabla_items_data = [[
        Paragraph("Items", header_table_style),
        Paragraph("Código / SKU", header_table_style),
        Paragraph("Producto / Descripción", header_table_style),
        Paragraph("Marca", header_table_style),
        Paragraph("Modelo", header_table_style),
        Paragraph("Cantidad", header_table_style)
    ]]

    total_cantidades = 0
    # Obtenemos los ítems ordenados por la sección del ítem de cotización asociado
    items = despacho.items.select_related('item_cotizacion__producto').order_by('item_cotizacion__seccion_departamento')
    
    # Agrupamos los ítems por sección para el ReportLab
    items_agrupados = []
    for seccion, grupo in groupby(items, key=lambda x: x.item_cotizacion.seccion_departamento):
        items_agrupados.append({'seccion': seccion, 'items': list(grupo)})

    for grupo in items_agrupados:
        # Encabezado de sección
        tabla_items_data.append([
            Paragraph(f"<b>{grupo['seccion']}</b>", bold_style), "", "", "", "", ""
        ])
        for idx, item in enumerate(grupo['items'], start=1):
            tabla_items_data.append([
                Paragraph(str(idx), normal_style),
                Paragraph(item.item_cotizacion.producto.codigo_sku, normal_style),
                Paragraph(item.item_cotizacion.producto.nombre, normal_style),
                Paragraph(item.marca or "-", normal_style),
                Paragraph(item.modelo or "-", normal_style),
                Paragraph(f"{item.cantidad_despachada:,.2f}", normal_style)
            ])
            total_cantidades += item.cantidad_despachada

    tabla_items_data.append([
        Paragraph("<b>TOTAL DE PRODUCTOS:</b>", bold_style), "", "", "", "",
        Paragraph(f"<b>{total_cantidades:,.2f}</b>", bold_style)
    ])

    table_items = Table(tabla_items_data, colWidths=[40, 90, 190, 75, 75, 80])
    table_items.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A252C')),
        ('GRID', (0, 0), (-1, -2), 0.5, colors.HexColor('#9CA3AF')),
        ('ALIGN', (5, 1), (5, -1), 'RIGHT'),
        ('SPAN', (0, -1), (4, -1)),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(table_items)
    story.append(Spacer(1, 15))

    data_transporte = [
        [Paragraph("<b>NOMBRE DEL CONDUCTOR:</b>", bold_style),
         Paragraph(despacho.nombre_conductor or "-", normal_style),
         Paragraph("<b>C.I. CONDUCTOR:</b>", bold_style),
         Paragraph(despacho.cedula_conductor or "-", normal_style)],
        [Paragraph("<b>TIPO DE VEHÍCULO:</b>", bold_style),
         Paragraph(despacho.tipo_vehiculo or "-", normal_style),
         Paragraph("<b>PLACA / COLOR:</b>", bold_style),
         Paragraph(f"{despacho.placa_vehiculo or '-'} / {despacho.color_vehiculo or '-'}", normal_style)],
        [Paragraph("<b>BENEFICIARIO AUTORIZADO:</b>", bold_style),
         Paragraph(despacho.beneficiario_autorizado or "-", normal_style),
         Paragraph("<b>C.I. / TELF:</b>", bold_style),
         Paragraph(f"{despacho.cedula_beneficiario or '-'} / {despacho.telefono_beneficiario or '-'}", normal_style)],
        [Paragraph("<b>LLEGADA CONDUCTOR:</b>", bold_style),
         Paragraph(despacho.fecha_hora_llegada_conductor.strftime('%d/%m/%Y %H:%M') if despacho.fecha_hora_llegada_conductor else "-", normal_style),
         Paragraph("<b>LLEGADA BENEFICIARIO:</b>", bold_style),
         Paragraph(despacho.fecha_hora_llegada_beneficiario.strftime('%d/%m/%Y %H:%M') if despacho.fecha_hora_llegada_beneficiario else "-", normal_style)]
    ]
    table_trans = Table(data_transporte, colWidths=[130, 145, 130, 145])
    table_trans.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor('#F9FAFB')),
        ('BACKGROUND', (2, 0), (2, -1), colors.HexColor('#F9FAFB')),
        ('PADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(table_trans)
    story.append(Spacer(1, 12))

    nota_legal = """<i><b>Nota importante:</b><br/>
    1.- La empresa Venezuela Productiva no se hace responsable de daños ocasionados a productos en el traslado del mismo.<br/>
    2.- Cualquier sustracción o deterioro que no conste en la presente guía, reclame al conductor de inmediato.<br/>
    3.- Favor devolver el duplicado de la guía firmado como garantía de recepción conforme.</i>
    """
    story.append(Paragraph(nota_legal, normal_style))
    story.append(Spacer(1, 35))

    data_firmas = [
        ["_________________________", "_________________________", "_________________________"],
        ["ELABORADO POR (Almacén)", "Conductor / Transportista", "Beneficiario / Autorizado"],
        [f"Usuario: {despacho.responsable.get_full_name() or despacho.responsable.username}",
         f"C.I.: {despacho.cedula_conductor or '-'}",
         f"C.I.: {despacho.cedula_beneficiario or '-'}"]
    ]
    table_firmas = Table(data_firmas, colWidths=[183, 183, 184])
    table_firmas.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 2),
    ]))
    story.append(table_firmas)

    doc.build(story)
    pdf = buffer.getvalue()
    buffer.close()
    response.write(pdf)
    return response

@login_required
def api_productos_cotizacion(request, cotizacion_id):
    cotizacion = get_object_or_404(Cotizacion, id=cotizacion_id)
    items = cotizacion.items.select_related('producto').all()
    data = []
    for item in items:
        data.append({
            'id_item_cotizacion': item.id,
            'producto_id': item.producto.id,
            'producto_nombre': f"[{item.producto.codigo_sku}] {item.producto.nombre}",
            'seccion': item.seccion_departamento,
            'cantidad_pendiente': float(item.cantidad_faltante),
        })
    return JsonResponse({'items': data})

# ==============================================================================
# VISTAS AUXILIARES Y CARGA MASIVA
# ==============================================================================

@login_required
def crear_cliente(request):
    if request.method == 'POST':
        form = ClienteForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('core:crear_cotizacion')
    else:
        form = ClienteForm()
    return render(request, 'core/cliente_form.html', {'form': form})

@login_required
def crear_proveedor(request):
    if request.method == 'POST':
        form = ProveedorForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('core:dashboard')
    else:
        form = ProveedorForm()
    return render(request, 'core/proveedor_form.html', {'form': form})

@login_required
def crear_almacen(request):
    if request.method == 'POST':
        form = AlmacenForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('core:dashboard')
    else:
        form = AlmacenForm()
    return render(request, 'core/almacen_form.html', {'form': form})

@login_required
def carga_masiva_productos(request):
    context = {'errores': [], 'exitos': 0}
    if request.method == 'POST' and request.FILES.get('archivo_xlsx'):
        archivo = request.FILES['archivo_xlsx']
        if not archivo.name.endswith('.xlsx'):
            context['errores'].append("El archivo debe ser formato .xlsx")
            return render(request, 'core/carga_masiva_form.html', context)

        try:
            wb = openpyxl.load_workbook(archivo, data_only=True)
            hoja = wb.active

            fila_encabezados = None
            num_fila_encabezados = 0
            for row in hoja.iter_rows(min_row=1, max_row=15, values_only=True):
                num_fila_encabezados += 1
                headers = [str(c).strip().upper() for c in row if c is not None]
                if 'CODIGO' in headers or 'DESCRIPCION' in headers:
                    fila_encabezados = [str(c).strip().upper() if c is not None else "" for c in row]
                    break

            if not fila_encabezados:
                context['errores'].append("No se encontró la fila de encabezados con 'CODIGO' o 'DESCRIPCION'.")
                return render(request, 'core/carga_masiva_form.html', context)

            col_map = {idx: name for idx, name in enumerate(fila_encabezados) if name}

            conteo_exitos = 0
            with transaction.atomic():
                for nro_fila, row in enumerate(
                        hoja.iter_rows(min_row=num_fila_encabezados + 1, values_only=True),
                        start=num_fila_encabezados + 1):

                    fila_dict = {}
                    for idx, val in enumerate(row):
                        header_name = col_map.get(idx)
                        if header_name:
                            fila_dict[header_name] = val

                    nombre = str(fila_dict.get('DESCRIPCION', '') or fila_dict.get('DESCRIPCION DEL PRODUCTO', '')).strip()
                    if not nombre or nombre == 'None' or "TOTAL" in nombre.upper():
                        continue

                    codigo = str(fila_dict.get('CODIGO', '')).strip()
                    stock_raw = fila_dict.get('INVENTARIO', fila_dict.get('CANTIDADES', fila_dict.get('ENTRADA', 0)))
                    codigo_oddo = str(fila_dict.get('CODIGO ODDO', '')).strip()
                    costo_raw = fila_dict.get('COSTO', fila_dict.get('COSTO_COMPRA', fila_dict.get('PRECIO UNITARIO', 0)))
                    venta_raw = fila_dict.get('PRECIO', fila_dict.get('PRECIO_VENTA', fila_dict.get('MONTO_VENTA', 0)))
                    almacen_raw = str(fila_dict.get('ALMACEN', '') or fila_dict.get('ALMACÉN', '')).strip()

                    sku_final = codigo if codigo and codigo != 'None' else codigo_oddo
                    if not sku_final or sku_final in ('None', ''):
                        sku_final = f"GEN-{nro_fila}"

                    def sanear_numero(val):
                        if val is None or str(val).strip() in ('', 'None'):
                            return 0.00
                        try:
                            return float(str(val).replace(',', '.'))
                        except ValueError:
                            return 0.00

                    stock_val = sanear_numero(stock_raw)
                    costo_val = sanear_numero(costo_raw)
                    venta_val = sanear_numero(venta_raw)

                    almacen_instancia = None
                    if almacen_raw and almacen_raw != 'None':
                        almacen_instancia, _ = Almacen.objects.get_or_create(
                            nombre=almacen_raw,
                            defaults={'ubicacion': 'Registrado vía Carga Masiva', 'activo': True}
                        )

                    Producto.objects.update_or_create(
                        codigo_sku=sku_final,
                        defaults={
                            'nombre': nombre,
                            'descripcion': f"Código Odoo: {codigo_oddo}" if codigo_oddo and codigo_oddo != 'None' else "Cargado masivamente",
                            'stock_actual': stock_val,
                            'costo_compra': costo_val,
                            'precio_venta': venta_val,
                            'almacen': almacen_instancia,
                            'activo': True
                        }
                    )
                    conteo_exitos += 1

            context['exitos'] = conteo_exitos

        except Exception as e:
            context['errores'].append(f"Error crítico procesando el archivo: {str(e)}")

    return render(request, 'core/carga_masiva_form.html', context)